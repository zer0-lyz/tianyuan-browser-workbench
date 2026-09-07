#!/usr/bin/env python3
"""Apply the Tianyuan table-format preset to one .docx document."""

from __future__ import annotations

import argparse
import json
import os
import posixpath
import re
import sys
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.table import _Cell
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from lxml import etree


FORMAT_RULES = {
    "fontEastAsia": "宋体",
    "fontLatin": "Times New Roman",
    "fontSizePt": 10,
    "fitToWindow": True,
    "minimumRowHeightCm": 0.6,
    "repeatHeader": True,
    "headerBold": True,
    "headerHorizontalAlignment": "center",
    "headerVerticalAlignment": "center",
    "numericAlignment": "right",
    "topBottomBorderPt": 1.5,
    "insideBorderPt": 0.5,
    "leftRightBorder": "none",
    "verticalAlignment": "center",
    "lineSpacing": 1.0,
}

DOCX_REQUIRED_MEMBERS = ("[Content_Types].xml", "_rels/.rels", "word/document.xml")

NUMERIC_CELL_RE = re.compile(
    r"^[\s\d,.，。;；:：'‘’\"“”%％+-−()（）/\\$￥€£¥*]+$"
)
ZERO_VALUE_RE = re.compile(r"^[+−-]?0(?:[.,，。]0+)?%?$")
REMARK_PLACEHOLDER_RE = re.compile(r"^(?:[-—–－]|[+−-]?0(?:[.,，。]0+)?%?)$")
SERIAL_HEADER_RE = re.compile(r"^(?:序号|序|项次|no\.?|no)$", re.IGNORECASE)
IDENTIFIER_HEADER_TERMS = (
    "编号", "编码", "代码", "证号", "文号", "统一社会信用代码", "社会信用代码",
)
SERIAL_VALUE_RE = re.compile(r"^(?:\d{1,4}|[A-Za-z]?\d{1,4})[.)、．。]?$")
HEADER_TERMS = (
    "序号", "编号", "名称", "项目", "内容", "金额", "原值", "净值", "价值", "数量", "日期", "时间",
    "标准", "评分", "描述", "类型", "状态", "情况", "比例", "权利", "申请", "注册", "单位", "位置",
    "面积", "年度", "方法", "成本", "价格", "层数", "结构", "备注", "公司", "科目", "用途", "权属",
    "类别", "持股", "建成", "年限", "成新率", "权重", "文号", "证号",
)

# A 10 pt amount with separators needs a little more room than the generic
# numeric floor. Keep the number at native size and reserve this width before
# allowing ordinary text columns to give up space.
NUMERIC_MIN_WIDTH_TWIPS = 900
NUMERIC_CELL_PADDING_TWIPS = 240
NUMERIC_WIDTH_PER_UNIT_TWIPS = 90
COMPACT_COLUMN_PADDING_TWIPS = 260
TEXT_MIN_WIDTH_TWIPS = 720
TEXT_MAX_MIN_WIDTH_TWIPS = 1000
SERIAL_MIN_WIDTH_TWIPS = 600
SERIAL_MAX_WIDTH_TWIPS = 1200
IDENTIFIER_MIN_WIDTH_TWIPS = 900
IDENTIFIER_MAX_WIDTH_TWIPS = 3000


def emit(event: str, **payload: object) -> None:
    message = {"event": event, **payload}
    print(json.dumps(message, ensure_ascii=False), flush=True)


def validate_docx_package(document_path: Path) -> None:
    """Fail before python-docx mutates a file that is not an OOXML package."""
    try:
        with zipfile.ZipFile(document_path, "r") as document_zip:
            names = set(document_zip.namelist())
            missing = [member for member in DOCX_REQUIRED_MEMBERS if member not in names]
            if missing:
                raise ValueError(
                    "TABLE_FORMAT_INVALID_DOCX_MISSING_MEMBER:" + ",".join(missing)
                )
            bad_member = document_zip.testzip()
            if bad_member:
                raise ValueError("TABLE_FORMAT_INVALID_DOCX_CORRUPT_MEMBER:" + bad_member)
    except zipfile.BadZipFile as error:
        raise ValueError("TABLE_FORMAT_INVALID_DOCX_NOT_ZIP") from error


def repair_null_relationships(document_path: Path) -> int:
    """Remove only package relationships that point to the literal missing NULL part."""
    removed_count = 0
    temporary_path = None
    try:
        with zipfile.ZipFile(document_path, "r") as source_zip:
            names = set(source_zip.namelist())
            relationship_updates = {}
            for name in names:
                if not name.endswith(".rels"):
                    continue
                relationship_xml = source_zip.read(name)
                root = etree.fromstring(relationship_xml)
                rel_directory = posixpath.dirname(name)
                base_directory = (
                    rel_directory[: -len("/_rels")]
                    if rel_directory.endswith("/_rels")
                    else ""
                )
                changed = False
                for relationship in list(root):
                    if relationship.get("TargetMode") == "External":
                        continue
                    target = str(relationship.get("Target") or "")
                    resolved_target = posixpath.normpath(
                        posixpath.join(base_directory, target)
                    ).lstrip("/")
                    if (
                        posixpath.basename(resolved_target).upper() == "NULL"
                        and resolved_target not in names
                    ):
                        root.remove(relationship)
                        removed_count += 1
                        changed = True
                if changed:
                    relationship_updates[name] = etree.tostring(
                        root,
                        xml_declaration=True,
                        encoding="UTF-8",
                        standalone=True,
                    )

            if not relationship_updates:
                return 0

            with tempfile.NamedTemporaryFile(
                prefix=f".{document_path.stem}.tianyuan-repair-",
                suffix=document_path.suffix,
                dir=document_path.parent,
                delete=False,
            ) as temporary:
                temporary_path = Path(temporary.name)
            with zipfile.ZipFile(temporary_path, "w", zipfile.ZIP_DEFLATED) as target_zip:
                for info in source_zip.infolist():
                    data = relationship_updates.get(info.filename, source_zip.read(info.filename))
                    target_zip.writestr(info, data)
            with zipfile.ZipFile(temporary_path, "r") as validation_zip:
                bad_member = validation_zip.testzip()
                if bad_member:
                    raise ValueError(f"TABLE_FORMAT_REPAIR_CORRUPT:{bad_member}")
            mode = document_path.stat().st_mode & 0o777
            os.replace(temporary_path, document_path)
            temporary_path = None
            os.chmod(document_path, mode)
            return removed_count
    finally:
        if temporary_path and temporary_path.exists():
            temporary_path.unlink()


def remove_children(parent, tag: str) -> None:
    for child in list(parent):
        if child.tag == qn(tag):
            parent.remove(child)


def set_run_fonts(run) -> None:
    run.font.name = FORMAT_RULES["fontLatin"]
    run.font.size = Pt(FORMAT_RULES["fontSizePt"])
    run_properties = run._element.get_or_add_rPr()
    fonts = run_properties.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        run_properties.insert(0, fonts)
    fonts.set(qn("w:ascii"), FORMAT_RULES["fontLatin"])
    fonts.set(qn("w:hAnsi"), FORMAT_RULES["fontLatin"])
    fonts.set(qn("w:eastAsia"), FORMAT_RULES["fontEastAsia"])


def is_numeric_cell(text: str) -> bool:
    normalized = re.sub(r"\s+", "", text or "")
    return bool(normalized and re.search(r"\d", normalized) and NUMERIC_CELL_RE.fullmatch(normalized))


def is_redundant_zero(text: str) -> bool:
    normalized = re.sub(r"\s+", "", text or "")
    return bool(ZERO_VALUE_RE.fullmatch(normalized))


def is_redundant_remark_value(text: str) -> bool:
    normalized = re.sub(r"\s+", "", text or "")
    return bool(REMARK_PLACEHOLDER_RE.fullmatch(normalized))


def cell_text_from_xml(cell) -> str:
    return "".join(node.text or "" for node in cell.iter(qn("w:t"))).strip()


def cell_grid_span(cell) -> int:
    cell_properties = cell.find(qn("w:tcPr"))
    if cell_properties is None:
        return 1
    grid_span = cell_properties.find(qn("w:gridSpan"))
    try:
        return max(1, int(grid_span.get(qn("w:val")))) if grid_span is not None else 1
    except (TypeError, ValueError):
        return 1


def row_grid_before(row) -> int:
    row_properties = row._tr.find(qn("w:trPr"))
    if row_properties is None:
        return 0
    grid_before = row_properties.find(qn("w:gridBefore"))
    try:
        return max(0, int(grid_before.get(qn("w:val")))) if grid_before is not None else 0
    except (TypeError, ValueError):
        return 0


def row_descriptors(row):
    column = row_grid_before(row)
    descriptors = []
    for cell in row._tr.iterchildren(qn("w:tc")):
        span = cell_grid_span(cell)
        descriptors.append({
            "cell": cell,
            "cell_object": _Cell(cell, row.table),
            "column": column,
            "span": span,
            "text": cell_text_from_xml(cell),
        })
        column += span
    return descriptors


def row_texts(row) -> list[str]:
    return [item["text"] for item in row_descriptors(row) if item["text"]]


def row_has_horizontal_merge(row) -> bool:
    return any(item["span"] > 1 for item in row_descriptors(row))


def normalized_header_text(text: str) -> str:
    return re.sub(r"[\s:：.．()（）\[\]【】]", "", text or "").lower()


def is_serial_header_text(text: str) -> bool:
    normalized = normalized_header_text(text)
    return bool(normalized and SERIAL_HEADER_RE.fullmatch(normalized))


def is_serial_value(text: str) -> bool:
    return bool(SERIAL_VALUE_RE.fullmatch(re.sub(r"\s+", "", text or "")))


def display_width(text: str) -> float:
    lines = str(text or "").splitlines() or [""]
    widths = []
    for line in lines:
        width = 0.0
        for char in line:
            if char.isspace():
                width += 0.5
            elif ord(char) >= 0x2E80:
                width += 2.0
            else:
                width += 1.0
        widths.append(width)
    return max(widths or [0.0])


def numeric_required_width(text: str) -> int:
    """Estimate the width needed for one-line numeric display at 10 pt."""
    normalized = re.sub(r"\s+", "", text or "")
    return max(
        NUMERIC_MIN_WIDTH_TWIPS,
        round(display_width(normalized) * NUMERIC_WIDTH_PER_UNIT_TWIPS + NUMERIC_CELL_PADDING_TWIPS),
    )


def compact_column_required_width(text: str, *, serial: bool) -> int:
    normalized = re.sub(r"\s+", "", text or "")
    estimated = round(
        display_width(normalized) * NUMERIC_WIDTH_PER_UNIT_TWIPS
        + COMPACT_COLUMN_PADDING_TWIPS
    )
    if serial:
        return max(SERIAL_MIN_WIDTH_TWIPS, min(SERIAL_MAX_WIDTH_TWIPS, estimated))
    return max(IDENTIFIER_MIN_WIDTH_TWIPS, min(IDENTIFIER_MAX_WIDTH_TWIPS, estimated))


def text_minimum_width(text: str) -> int:
    normalized = re.sub(r"\s+", "", text or "")
    estimated = round(
        display_width(normalized) * NUMERIC_WIDTH_PER_UNIT_TWIPS
        + COMPACT_COLUMN_PADDING_TWIPS
    )
    return max(TEXT_MIN_WIDTH_TWIPS, min(TEXT_MAX_MIN_WIDTH_TWIPS, estimated))


def is_data_like_text(text: str) -> bool:
    normalized = re.sub(r"\s+", "", text or "")
    if not normalized:
        return False
    if re.fullmatch(r"[-+]?\d[\d,.，]*(?:\.\d+)?%?", normalized):
        return True
    if re.search(r"\d", normalized) and (
        re.search(r"\d[.,/]\d", normalized)
        or "%" in normalized
        or "％" in normalized
        or ("年" in normalized and "月" in normalized)
    ):
        return True
    return display_width(normalized) > 15


def is_strong_value_text(text: str) -> bool:
    """Identify values in a label/value metadata row without treating group headers as values."""
    normalized = re.sub(r"\s+", "", text or "")
    if not normalized:
        return False
    return bool(
        is_data_like_text(normalized)
        or re.search(r"\d", normalized)
        or display_width(normalized) > 10
    )


def is_context_row(row) -> bool:
    """Exclude merged label/value metadata rows from repeated table headers."""
    descriptors = [item for item in row_descriptors(row) if item["text"]]
    if len(descriptors) < 3 or sum(item["span"] > 1 for item in descriptors) < 2:
        return False
    values = descriptors[1::2]
    labels = descriptors[::2]
    strong_values = sum(1 for item in values if is_strong_value_text(item["text"]))
    label_hits = sum(
        1 for item in labels
        if any(term in item["text"] for term in HEADER_TERMS)
    )
    return strong_values >= max(1, (len(values) + 1) // 2) and label_hits >= 1


def header_score(row) -> tuple[float, float]:
    texts = row_texts(row)
    if len(texts) < 2:
        return 0.0, 0.0
    header_hits = sum(1 for text in texts if any(term in text for term in HEADER_TERMS))
    data_like = sum(1 for text in texts if is_data_like_text(text))
    return header_hits / len(texts), data_like / len(texts)


def is_header_like(row) -> bool:
    if is_context_row(row):
        return False
    score, data_like = header_score(row)
    return score >= 0.5 and data_like <= 0.25


def is_group_header_row(row) -> bool:
    """Recognize a merged grouping row immediately above a normal header row."""
    if is_context_row(row):
        return False
    texts = row_texts(row)
    if not texts or not row_has_horizontal_merge(row):
        return False
    if any(is_strong_value_text(text) for text in texts):
        return False
    score, data_like = header_score(row)
    return data_like == 0.0 and (len(texts) <= 2 or score >= 0.5)


def detect_header_rows(table) -> set[int]:
    rows = table.rows
    if not rows:
        return set()
    candidate = None
    for index, row in enumerate(rows[:6]):
        if is_header_like(row):
            candidate = index
            break
    if candidate is None:
        return {0}
    header_rows = {candidate}
    if candidate > 0 and is_group_header_row(rows[candidate - 1]):
        header_rows.add(candidate - 1)
    if candidate == 0 and len(rows) > 1 and is_header_like(rows[1]):
        header_rows.add(1)
    return header_rows


def column_has_serial_values(table, column: int, header_rows: set[int]) -> bool:
    values = []
    for row_index, row in enumerate(table.rows):
        if row_index in header_rows:
            continue
        for descriptor in row_descriptors(row):
            if descriptor["column"] == column and descriptor["span"] == 1 and descriptor["text"]:
                values.append(descriptor["text"])
    if len(values) < 2:
        return False
    matches = sum(1 for value in values[:12] if is_serial_value(value))
    return matches >= 2 and matches / min(len(values), 12) >= 0.75


def detect_serial_columns(table, header_rows: set[int]) -> set[int]:
    """Find the logical grid columns headed by 序号, with a conservative value fallback."""
    for row_index in sorted(header_rows):
        for descriptor in row_descriptors(table.rows[row_index]):
            if is_serial_header_text(descriptor["text"]):
                return set(range(
                    descriptor["column"],
                    descriptor["column"] + descriptor["span"],
                ))

    # Some templates use "编号" or omit the header text. Only infer it when
    # the values are visibly ordinal, so amounts and years are not centered.
    candidates = sorted({
        descriptor["column"]
        for row in table.rows
        for descriptor in row_descriptors(row)
        if descriptor["span"] == 1
    })
    for column in candidates[:4]:
        if column_has_serial_values(table, column, header_rows):
            return {column}
    return set()


def is_identifier_header_text(text: str) -> bool:
    normalized = normalized_header_text(text)
    return bool(normalized and any(term in normalized for term in IDENTIFIER_HEADER_TERMS))


def detect_identifier_columns(table, header_rows: set[int]) -> set[int]:
    columns = set()
    for row_index in sorted(header_rows):
        for descriptor in row_descriptors(table.rows[row_index]):
            if is_identifier_header_text(descriptor["text"]):
                columns.update(range(
                    descriptor["column"],
                    descriptor["column"] + descriptor["span"],
                ))
    return columns


def column_content_width(table, column: int) -> float:
    values = [
        display_width(descriptor["text"])
        for row in table.rows
        for descriptor in row_descriptors(row)
        if descriptor["column"] == column
        and descriptor["span"] == 1
        and descriptor["text"]
    ]
    return max(values or [0.0])


def clear_redundant_remark_values(table, header_rows: set[int]) -> int:
    remark_columns = set()
    for row_index in sorted(header_rows):
        for descriptor in row_descriptors(table.rows[row_index]):
            normalized = normalized_header_text(descriptor["text"])
            if "备注" in normalized or "说明" in normalized:
                remark_columns.update(range(
                    descriptor["column"],
                    descriptor["column"] + descriptor["span"],
                ))
    if not remark_columns:
        return 0
    cleaned = 0
    for row_index, row in enumerate(table.rows):
        if row_index in header_rows:
            continue
        for descriptor in row_descriptors(row):
            if (
                descriptor["span"] == 1
                and descriptor["column"] in remark_columns
                and is_redundant_remark_value(descriptor["text"])
            ):
                descriptor["cell_object"].text = ""
                cleaned += 1
    return cleaned


def set_cell_vertical_alignment(cell, value: WD_CELL_VERTICAL_ALIGNMENT) -> None:
    cell.vertical_alignment = value
    cell_properties = cell._tc.get_or_add_tcPr()
    v_align = cell_properties.find(qn("w:vAlign"))
    if v_align is None:
        v_align = OxmlElement("w:vAlign")
        cell_properties.append(v_align)
    v_align.set(qn("w:val"), "center" if value == WD_CELL_VERTICAL_ALIGNMENT.CENTER else "top")


def clear_cell_fit_text(cell) -> None:
    cell_properties = cell._tc.get_or_add_tcPr()
    remove_children(cell_properties, "w:tcFitText")


def set_paragraph_alignment(paragraph, alignment) -> None:
    """Replace inherited justify/distribute alignment with an explicit value."""
    paragraph.alignment = alignment
    paragraph_properties = paragraph._p.get_or_add_pPr()
    remove_children(paragraph_properties, "w:jc")
    value = "left"
    if alignment == WD_ALIGN_PARAGRAPH.CENTER:
        value = "center"
    elif alignment == WD_ALIGN_PARAGRAPH.RIGHT:
        value = "right"
    justification = OxmlElement("w:jc")
    justification.set(qn("w:val"), value)
    paragraph_properties.append(justification)


def set_single_line_spacing(paragraph) -> None:
    paragraph.paragraph_format.line_spacing = FORMAT_RULES["lineSpacing"]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph_properties = paragraph._p.get_or_add_pPr()
    spacing = paragraph_properties.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        paragraph_properties.append(spacing)
    spacing.set(qn("w:before"), "0")
    spacing.set(qn("w:after"), "0")
    spacing.set(qn("w:line"), "240")
    spacing.set(qn("w:lineRule"), "auto")


def set_cell_text_format(
    cell,
    *,
    header: bool,
    serial: bool = False,
    compact: bool = False,
    preserve_no_wrap: bool = True,
) -> None:
    text = "\n".join(paragraph.text for paragraph in cell.paragraphs).strip()
    numeric = is_numeric_cell(text)
    alignment = WD_ALIGN_PARAGRAPH.CENTER if header or serial else (
        WD_ALIGN_PARAGRAPH.RIGHT if numeric else WD_ALIGN_PARAGRAPH.LEFT
    )
    cell_properties = cell._tc.get_or_add_tcPr()
    remove_children(cell_properties, "w:noWrap")
    if preserve_no_wrap and (numeric or serial or compact):
        cell_properties.append(OxmlElement("w:noWrap"))
    # Never shrink numeric glyphs to force them into a narrow column. The
    # preset keeps the original 10 pt scale and lets the width calculation
    # decide how much space the column should receive.
    clear_cell_fit_text(cell)
    for paragraph in cell.paragraphs:
        set_paragraph_alignment(paragraph, alignment)
        set_single_line_spacing(paragraph)
        for run in paragraph.runs:
            set_run_fonts(run)
            if header:
                run.bold = True
    set_cell_vertical_alignment(cell, WD_CELL_VERTICAL_ALIGNMENT.CENTER)


def set_row_header(row, enabled: bool) -> None:
    row_properties = row._tr.get_or_add_trPr()
    remove_children(row_properties, "w:tblHeader")
    if enabled:
        row_properties.append(OxmlElement("w:tblHeader"))


def border_size_eighths(points: float) -> str:
    return str(round(points * 8))


def set_table_borders(table) -> None:
    table_properties = table._tbl.tblPr
    borders = table_properties.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        table_properties.append(borders)
    remove_children(borders, "w:top")
    remove_children(borders, "w:bottom")
    remove_children(borders, "w:left")
    remove_children(borders, "w:right")
    remove_children(borders, "w:start")
    remove_children(borders, "w:end")
    remove_children(borders, "w:insideH")
    remove_children(borders, "w:insideV")

    def add_border(name: str, value: str, size: str | None = None) -> None:
        border = OxmlElement(f"w:{name}")
        border.set(qn("w:val"), value)
        if size:
            border.set(qn("w:sz"), size)
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "000000")
        borders.append(border)

    add_border("top", "single", border_size_eighths(FORMAT_RULES["topBottomBorderPt"]))
    add_border("bottom", "single", border_size_eighths(FORMAT_RULES["topBottomBorderPt"]))
    add_border("left", "nil")
    add_border("right", "nil")
    add_border("insideH", "single", border_size_eighths(FORMAT_RULES["insideBorderPt"]))
    add_border("insideV", "single", border_size_eighths(FORMAT_RULES["insideBorderPt"]))


def cell_vertical_merge_state(cell) -> str | None:
    cell_properties = cell.find(qn("w:tcPr"))
    if cell_properties is None:
        return None
    vertical_merge = cell_properties.find(qn("w:vMerge"))
    if vertical_merge is None:
        return None
    return vertical_merge.get(qn("w:val")) or "continue"


def find_descriptor(row, column: int, span: int):
    for descriptor in row_descriptors(row):
        if descriptor["column"] == column and descriptor["span"] == span:
            return descriptor
    return None


def add_cell_border(border_container, name: str, value: str, size: str | None = None) -> None:
    border = OxmlElement(f"w:{name}")
    border.set(qn("w:val"), value)
    if size:
        border.set(qn("w:sz"), size)
    border.set(qn("w:space"), "0")
    border.set(qn("w:color"), "000000")
    border_container.append(border)


def set_cell_border(
    cell,
    *,
    top: tuple[str, str | None],
    bottom: tuple[str, str | None],
    left: tuple[str, str | None],
    right: tuple[str, str | None],
) -> None:
    cell_properties = cell.find(qn("w:tcPr"))
    if cell_properties is None:
        cell_properties = OxmlElement("w:tcPr")
        cell.insert(0, cell_properties)
    borders = cell_properties.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        cell_properties.append(borders)
    for child in list(borders):
        borders.remove(child)
    add_cell_border(borders, "top", top[0], top[1])
    add_cell_border(borders, "bottom", bottom[0], bottom[1])
    add_cell_border(borders, "left", left[0], left[1])
    add_cell_border(borders, "right", right[0], right[1])


def set_table_cell_borders(table, column_count: int) -> None:
    rows = list(table.rows)
    thin = border_size_eighths(FORMAT_RULES["insideBorderPt"])
    thick = border_size_eighths(FORMAT_RULES["topBottomBorderPt"])
    for row_index, row in enumerate(rows):
        # Existing documents may carry row-level border exceptions that override
        # tblBorders. Remove only that exception and retain any row margins.
        for row_table_properties in row._tr.iter(qn("w:tblPrEx")):
            remove_children(row_table_properties, "w:tblBorders")
            if len(row_table_properties) == 0:
                parent = row_table_properties.getparent()
                if parent is not None:
                    parent.remove(row_table_properties)
        for descriptor in row_descriptors(row):
            start = descriptor["column"]
            end = start + descriptor["span"]
            merge_state = cell_vertical_merge_state(descriptor["cell"])
            next_descriptor = (
                find_descriptor(rows[row_index + 1], start, descriptor["span"])
                if row_index + 1 < len(rows)
                else None
            )
            continues_below = bool(
                next_descriptor
                and cell_vertical_merge_state(next_descriptor["cell"]) == "continue"
            )
            top_border = (
                ("nil", None)
                if merge_state == "continue"
                else ("single", thick if row_index == 0 else thin)
            )
            bottom_border = (
                ("nil", None)
                if continues_below
                else ("single", thick if row_index == len(rows) - 1 else thin)
            )
            left_border = ("nil", None) if start == 0 else ("single", thin)
            right_border = ("nil", None) if end >= column_count else ("single", thin)
            set_cell_border(
                descriptor["cell"],
                top=top_border,
                bottom=bottom_border,
                left=left_border,
                right=right_border,
            )


def document_table_width_twips(document) -> int:
    try:
        section = document.sections[0]
        page_width = int(section.page_width) // 635
        margins = (int(section.left_margin) + int(section.right_margin)) // 635
        return max(6000, min(12000, page_width - margins))
    except (AttributeError, IndexError, TypeError, ValueError):
        return 9000


def table_available_width_twips(table, document_width_twips: int) -> int:
    """Use the parent cell width for nested tables instead of page width."""
    parent = table._tbl.getparent()
    if parent is None or parent.tag != qn("w:tc"):
        return document_width_twips
    cell_properties = parent.find(qn("w:tcPr"))
    cell_width = cell_properties.find(qn("w:tcW")) if cell_properties is not None else None
    if cell_width is None:
        return document_width_twips
    try:
        value = int(cell_width.get(qn("w:w")))
    except (TypeError, ValueError):
        return document_width_twips
    if cell_width.get(qn("w:type")) == "pct":
        value = round(document_width_twips * value / 5000)
    return max(720, min(document_width_twips, value))


def set_table_width(table, width_twips: int) -> list:
    table.autofit = True
    table_properties = table._tbl.tblPr
    table_width = table_properties.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        table_properties.append(table_width)
    table_width.set(qn("w:w"), "5000")
    table_width.set(qn("w:type"), "pct")
    layout = table_properties.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        table_properties.append(layout)
    layout.set(qn("w:type"), "autofit")
    table_indent = table_properties.find(qn("w:tblInd"))
    if table_indent is None:
        table_indent = OxmlElement("w:tblInd")
        table_properties.append(table_indent)
    table_indent.set(qn("w:w"), "0")
    table_indent.set(qn("w:type"), "dxa")
    table_grid = table._tbl.tblGrid
    if table_grid is None:
        table_grid = OxmlElement("w:tblGrid")
        table._tbl.insert(1, table_grid)
    grid_columns = list(table_grid.iterchildren(qn("w:gridCol")))
    column_count = max(len(grid_columns), len(table.columns))
    while len(grid_columns) < column_count:
        grid_column = OxmlElement("w:gridCol")
        table_grid.append(grid_column)
        grid_columns.append(grid_column)
    return grid_columns


def calculate_column_widths(table, width_twips: int, header_rows: set[int]) -> tuple[list[int], bool]:
    grid_columns = list(table._tbl.tblGrid.iterchildren(qn("w:gridCol"))) if table._tbl.tblGrid is not None else []
    column_count = max(len(grid_columns), len(table.columns))
    if column_count <= 0:
        return [], False
    serial_columns = detect_serial_columns(table, header_rows)
    identifier_columns = detect_identifier_columns(table, header_rows)
    original_widths = []
    for grid_column in grid_columns[:column_count]:
        try:
            original_widths.append(max(360, int(grid_column.get(qn("w:w")))) )
        except (TypeError, ValueError):
            original_widths.append(0)
    fallback_width = max(360, round(width_twips / column_count))
    while len(original_widths) < column_count:
        original_widths.append(fallback_width)
    # Use the document's existing grid as the starting point. It usually
    # contains deliberate proportions that are lost when every column starts
    # from the same generic weight.
    weights = [width or fallback_width for width in original_widths]
    minimums = [360] * column_count
    for row in table.rows:
        for descriptor in row_descriptors(row):
            start = descriptor["column"]
            span = descriptor["span"]
            if start >= column_count:
                continue
            span = min(span, column_count - start)
            if is_numeric_cell(descriptor["text"]):
                desired = numeric_required_width(descriptor["text"])
                if span == 1:
                    minimums[start] = max(minimums[start], desired)
            elif start in serial_columns:
                desired = compact_column_required_width(descriptor["text"], serial=True)
                if span == 1:
                    minimums[start] = max(minimums[start], SERIAL_MIN_WIDTH_TWIPS)
            elif start in identifier_columns:
                desired = compact_column_required_width(descriptor["text"], serial=False)
                if span == 1:
                    minimums[start] = max(minimums[start], IDENTIFIER_MIN_WIDTH_TWIPS)
            else:
                desired = min(
                    5000,
                    max(600, round(display_width(descriptor["text"]) * NUMERIC_WIDTH_PER_UNIT_TWIPS + COMPACT_COLUMN_PADDING_TWIPS)),
                )
                if span == 1:
                    minimums[start] = max(minimums[start], text_minimum_width(descriptor["text"]))
            if span == 1:
                weights[start] = max(weights[start], desired)
            else:
                current = sum(weights[start : start + span])
                if desired > current:
                    extra = (desired - current) / span
                    for index in range(start, start + span):
                        weights[index] += extra

    # Compact columns are sized from their own content rather than inheriting
    # the generic text/numeric minimum. This keeps 序号 narrow while allowing
    # a real 编号 value to remain complete.
    for column in sorted(serial_columns | identifier_columns):
        if column >= column_count:
            continue
        is_serial = column in serial_columns
        column_values = [
            descriptor["text"]
            for row in table.rows
            for descriptor in row_descriptors(row)
            if descriptor["column"] == column
            and descriptor["span"] == 1
            and descriptor["text"]
        ]
        required = max(
            (
                compact_column_required_width(value, serial=is_serial)
                for value in column_values
            ),
            default=SERIAL_MIN_WIDTH_TWIPS if is_serial else IDENTIFIER_MIN_WIDTH_TWIPS,
        )
        content_width = column_content_width(table, column)
        weights[column] = max(
            required,
            round(content_width * NUMERIC_WIDTH_PER_UNIT_TWIPS + COMPACT_COLUMN_PADDING_TWIPS),
        )
        minimums[column] = required
    widths = [max(minimums[index], round(width_twips * weight / sum(weights))) for index, weight in enumerate(weights)]
    difference = width_twips - sum(widths)
    if difference:
        order = sorted(
            range(column_count),
            key=lambda index: (
                minimums[index] != 360 if difference < 0 else False,
                weights[index],
            ),
            reverse=difference > 0,
        )
        step = 1 if difference > 0 else -1
        remaining = abs(difference)
        while remaining:
            changed = False
            for index in order:
                if not remaining:
                    break
                if step > 0 or widths[index] > minimums[index]:
                    widths[index] += step
                    remaining -= 1
                    changed = True
            if not changed:
                break
    return fit_column_widths_to_window(widths, width_twips)


def fit_column_widths_to_window(widths: list[int], width_twips: int) -> tuple[list[int], bool]:
    total = sum(widths)
    if not widths or total <= width_twips:
        return widths, False
    minimum = min(240, max(120, width_twips // len(widths)))
    scaled = [max(minimum, round(width * width_twips / total)) for width in widths]
    while sum(scaled) > width_twips:
        candidates = [index for index, value in enumerate(scaled) if value > minimum]
        if not candidates:
            break
        index = max(candidates, key=scaled.__getitem__)
        scaled[index] -= 1
    index = 0
    while sum(scaled) < width_twips:
        scaled[index % len(scaled)] += 1
        index += 1
    return scaled, True


def set_column_widths(table, grid_columns: list, widths: list[int]) -> None:
    for grid_column, width in zip(grid_columns, widths):
        grid_column.set(qn("w:w"), str(width))
    for row in table.rows:
        for descriptor in row_descriptors(row):
            start = descriptor["column"]
            span = min(descriptor["span"], max(0, len(widths) - start))
            if span <= 0:
                continue
            cell_properties = descriptor["cell"].find(qn("w:tcPr"))
            if cell_properties is None:
                cell_properties = OxmlElement("w:tcPr")
                descriptor["cell"].insert(0, cell_properties)
            cell_width = cell_properties.find(qn("w:tcW"))
            if cell_width is None:
                cell_width = OxmlElement("w:tcW")
                cell_properties.insert(0, cell_width)
            cell_width.set(qn("w:w"), str(sum(widths[start : start + span])))
            cell_width.set(qn("w:type"), "dxa")


def set_table_rows(table, header_rows: set[int], *, preserve_no_wrap: bool = True) -> None:
    serial_columns = detect_serial_columns(table, header_rows)
    identifier_columns = detect_identifier_columns(table, header_rows)
    for row_index, row in enumerate(table.rows):
        row.height = Cm(FORMAT_RULES["minimumRowHeightCm"])
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        is_header = FORMAT_RULES["repeatHeader"] and row_index in header_rows
        set_row_header(row, is_header)
        for descriptor in row_descriptors(row):
            serial = (
                descriptor["span"] == 1
                and descriptor["column"] in serial_columns
            )
            identifier = (
                descriptor["span"] == 1
                and descriptor["column"] in identifier_columns
            )
            compact = serial or identifier
            set_cell_text_format(
                descriptor["cell_object"],
                header=is_header,
                serial=serial,
                compact=compact,
                preserve_no_wrap=preserve_no_wrap,
            )


def iter_tables(tables):
    for table in tables:
        yield table
        for row in table.rows:
            for cell in row.cells:
                yield from iter_tables(cell.tables)


def format_document(document: Document) -> tuple[int, int]:
    table_list = list(iter_tables(document.tables))
    document_width_twips = document_table_width_twips(document)
    cleaned_remark_values = 0
    for table in table_list:
        width_twips = table_available_width_twips(table, document_width_twips)
        header_rows = detect_header_rows(table)
        cleaned_remark_values += clear_redundant_remark_values(table, header_rows)
        grid_columns = set_table_width(table, width_twips)
        widths, compressed = calculate_column_widths(table, width_twips, header_rows)
        set_column_widths(table, grid_columns, widths)
        set_table_borders(table)
        set_table_rows(table, header_rows, preserve_no_wrap=not compressed)
        set_table_cell_borders(table, len(grid_columns))
    return len(table_list), cleaned_remark_values


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("document", type=Path)
    parser.add_argument("--no-backup", action="store_true", help=argparse.SUPPRESS)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    document_path = args.document.expanduser().resolve()
    if document_path.suffix.lower() != ".docx":
        raise ValueError("TABLE_FORMAT_DOCX_ONLY")
    if not document_path.is_file():
        raise FileNotFoundError("TABLE_FORMAT_INPUT_NOT_FOUND")
    validate_docx_package(document_path)
    emit("start", fileName=document_path.name)
    repaired_relationships = repair_null_relationships(document_path)
    if repaired_relationships:
        emit(
            "repaired",
            fileName=document_path.name,
            repairedRelationships=repaired_relationships,
        )
    document = Document(str(document_path))
    table_count, cleaned_remark_values = format_document(document)
    emit(
        "formatted",
        fileName=document_path.name,
        tableCount=table_count,
        repairedRelationships=repaired_relationships,
        cleanedRemarkValues=cleaned_remark_values,
    )
    document.save(str(document_path))
    emit(
        "saved",
        fileName=document_path.name,
        tableCount=table_count,
        repairedRelationships=repaired_relationships,
        cleanedRemarkValues=cleaned_remark_values,
    )
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as error:  # Keep the native helper error readable and secret-free.
        print(f"TABLE_FORMAT_FAILED:{error}", file=sys.stderr, flush=True)
        raise SystemExit(1)
