from __future__ import annotations

import hashlib
import re
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


SKILL_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SKILL_ROOT / "scripts"))

from depreciation_forecast.model import AssetForecast, ForecastParams  # noqa: E402
from depreciation_forecast.workbook import (  # noqa: E402
    ADDED_SHEET,
    PARAM_SHEET,
    STOCK_HEADERS,
    STOCK_SHEET,
    _output_totals,
    create_template,
    read_input,
)
from workflow import run_workflow  # noqa: E402


class WorkbookInputTests(unittest.TestCase):
    def _workbook(self, remaining_formula: str) -> Path:
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        path = Path(directory.name) / "input.xlsx"
        create_template(path)
        wb = load_workbook(path)
        wb[PARAM_SHEET]["B4"] = date(2024, 5, 31)
        wb[PARAM_SHEET]["B5"] = date(2029, 12, 31)
        stock = wb[STOCK_SHEET]
        stock.insert_cols(5, 2)
        stock.cell(3, 5, "资产科目1")
        stock.cell(3, 6, "设备编号")
        columns = {stock.cell(3, col).value: col for col in range(1, stock.max_column + 1)}
        if remaining_formula in {"MIN", "MAX"}:
            dep = f"{get_column_letter(columns['折旧年限'])}4"
            economic = f"{get_column_letter(columns['经济耐用年限'])}4"
            service = f"{get_column_letter(columns['启用时间'])}4"
            remaining_formula = f"=MAX({remaining_formula}({dep},{economic})-(参数设定!$B$4-{service})/365,0)"
        values = {
            "序号": "366",
            "情形描述": "测试",
            "公司主体": "中显",
            "资产科目": "机器设备",
            "名称": "测试资产",
            "账面原值": 100,
            "账面净值": 80,
            "评估原值": 100,
            "启用时间": date(2022, 5, 31),
            "折旧年限": 4,
            "经济耐用年限": 10,
            "预计尚可使用年限": remaining_formula,
            "更新后折旧年限": f"={get_column_letter(columns['折旧年限'])}4",
            "更新后经济耐用年限": f"={get_column_letter(columns['经济耐用年限'])}4",
            "残值率": 0,
            "费用科目": "主营业务成本",
            "折旧摊销": "折旧",
            "进项税率": 0.13,
            "是否需要更新": "是",
        }
        for header, value in values.items():
            stock.cell(4, columns[header], value)
        stock.cell(4, columns["资产科目1"], "元数据")
        stock.cell(4, columns["设备编号"], "E-001")
        wb.save(path)
        return path

    def test_formula_cells_work_without_cached_values(self) -> None:
        params, assets, _, errors = read_input(
            self._workbook("MAX")
        )
        self.assertFalse(errors)
        self.assertEqual(params.valuation_date, date(2024, 5, 31))
        self.assertEqual(len(assets), 1)
        self.assertAlmostEqual(assets[0].expected_remaining_years, 10 - 731 / 365)
        self.assertEqual(assets[0].renewed_depreciation_years, 4)
        self.assertEqual(assets[0].renewed_economic_years, 10)

    def test_min_variant_uses_the_formula_operator(self) -> None:
        _, assets, _, errors = read_input(
            self._workbook("MIN")
        )
        self.assertFalse(errors)
        self.assertAlmostEqual(assets[0].expected_remaining_years, 4 - 731 / 365)

    def test_unsupported_numeric_formula_fails_validation(self) -> None:
        _, _, _, errors = read_input(self._workbook("=SUM(J4:K4)"))
        self.assertTrue(any("未被支持" in error for error in errors))

    def test_custom_workbook_path_is_run_in_place(self) -> None:
        path = self._workbook("MAX")
        payload = run_workflow(workbook_path=path)
        self.assertEqual(payload["status"], "ok")
        self.assertEqual(Path(payload["workbook_path"]), path.resolve())


class PerpetualOutputTests(unittest.TestCase):
    def test_perpetual_values_are_not_reannualized_on_export(self) -> None:
        params = ForecastParams(date(2024, 5, 31), date(2029, 12, 31), 0.10)
        asset = AssetForecast(
            asset_kind="存量",
            asset_id="1",
            company="中显",
            account="机器设备",
            name="测试资产",
            expense_account="主营业务成本",
            scenario="",
            monthly={},
            perpetual={"折旧摊销": 100.0, "更新资本性支出": 25.0},
        )
        _, perpetual, expense = _output_totals([asset], params, params.end_date)
        self.assertEqual(perpetual["折旧摊销"], 100.0)
        self.assertEqual(perpetual["更新资本性支出"], 25.0)
        self.assertEqual(perpetual["资本性支出"], 25.0)
        self.assertEqual(expense["主营业务成本", "永续期"], 100.0)


class BundleManifestTests(unittest.TestCase):
    def test_manifest_lists_required_files_and_matches_digest(self) -> None:
        required_files = [
            "SKILL.md",
            "agents/openai.yaml",
            "references/workflow.md",
            "assets/折旧摊销预测输入模板.xlsx",
            "scripts/workflow.py",
            "scripts/depreciation_forecast/__init__.py",
            "scripts/depreciation_forecast/__main__.py",
            "scripts/depreciation_forecast/cli.py",
            "scripts/depreciation_forecast/model.py",
            "scripts/depreciation_forecast/workbook.py",
            "tests/test_workbook.py",
        ]
        manifest = (SKILL_ROOT / "agents" / "openai.yaml").read_text(encoding="utf-8")
        for relative_path in required_files:
            self.assertIn(f"- {relative_path}", manifest)
            self.assertTrue((SKILL_ROOT / relative_path).is_file())

        expected = re.search(r'value:\s*"([0-9a-f]{64})"', manifest)
        self.assertIsNotNone(expected)
        digest = hashlib.sha256()
        for relative_path in required_files:
            if relative_path == "agents/openai.yaml":
                continue
            digest.update(relative_path.encode("utf-8"))
            digest.update(b"\0")
            digest.update((SKILL_ROOT / relative_path).read_bytes())
            digest.update(b"\0")
        self.assertEqual(expected.group(1), digest.hexdigest())


if __name__ == "__main__":
    unittest.main()
