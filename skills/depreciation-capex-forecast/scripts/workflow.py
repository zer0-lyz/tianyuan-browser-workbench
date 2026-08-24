#!/usr/bin/env python3
from __future__ import annotations

import argparse
import hashlib
import json
import os
import shutil
import subprocess
import sys
import tempfile
import zipfile
from datetime import date, datetime
from pathlib import Path
from typing import Any, Callable

from openpyxl import load_workbook


SKILL_ROOT = Path(__file__).resolve().parents[1]
ENGINE_ROOT = SKILL_ROOT / "scripts"
ASSETS_ROOT = SKILL_ROOT / "assets"
CANONICAL_TEMPLATE = ASSETS_ROOT / "折旧摊销预测输入模板.xlsx"
WORKBOOK_PATH = Path.home() / "Downloads" / "折旧摊销预测输入模板.xlsx"

PARAM_SHEET = "参数设定"
STOCK_SHEET = "存量资产输入"
ADDED_SHEET = "新增资产输入"
SUMMARY_SHEET = "预测汇总"
EXPENSE_SHEET = "费用科目汇总"
CHECKS_SHEET = "检查结果"
ANNUAL_DETAIL_SHEET = "资产结果明细（年度）"
MONTHLY_DETAIL_SHEET = "资产结果明细（月度）"
DETAIL_PROCESS_SHEET = "详细过程"

STOCK_HEADERS = [
    "序号", "情形描述", "公司主体", "资产科目", "名称", "账面原值", "账面净值",
    "评估原值", "启用时间", "折旧年限", "经济耐用年限", "预计尚可使用年限",
    "更新后折旧年限", "更新后经济耐用年限", "残值率", "费用科目", "折旧摊销",
    "进项税率", "是否需要更新",
]
ADDED_HEADERS = [
    "序号", "情形描述", "公司主体", "资产科目", "名称", "在建工程账面价值",
    "总投资额（不含税）", "预计投入使用时间", "更新后折旧年限", "更新后经济耐用年限",
    "残值率", "费用科目", "折旧摊销", "进项税率", "是否需要更新",
]
PARAMETER_ROWS = {
    "评估基准日": 4,
    "结束日期": 5,
    "折现率": 6,
    "最低尚可使用年限": 7,
}
REQUIRED_HEADERS = {
    STOCK_SHEET: STOCK_HEADERS,
    ADDED_SHEET: ADDED_HEADERS,
}
RESULT_SHEETS = {
    SUMMARY_SHEET,
    CHECKS_SHEET,
    ANNUAL_DETAIL_SHEET,
    MONTHLY_DETAIL_SHEET,
    DETAIL_PROCESS_SHEET,
}

if str(ENGINE_ROOT) not in sys.path:
    sys.path.insert(0, str(ENGINE_ROOT))
from depreciation_forecast.workbook import read_input  # noqa: E402


def _python_env() -> dict[str, str]:
    env = os.environ.copy()
    existing = env.get("PYTHONPATH")
    pythonpath = str(ENGINE_ROOT)
    env["PYTHONPATH"] = f"{pythonpath}{os.pathsep}{existing}" if existing else pythonpath
    return env


def _resolved_path(value: str | Path) -> Path:
    return Path(value).expanduser().resolve()


def _ensure_workbook_path(value: str | Path) -> Path:
    path = _resolved_path(value)
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("WORKBOOK_EXTENSION_NOT_SUPPORTED")
    return path


def _ensure_existing_workbook(value: str | Path) -> Path:
    path = _ensure_workbook_path(value)
    if not path.exists():
        raise FileNotFoundError(f"未找到指定输入工作簿：{path}")
    if not path.is_file():
        raise ValueError("WORKBOOK_PATH_NOT_FILE")
    return path


def _temporary_path(target: Path) -> Path:
    target.parent.mkdir(parents=True, exist_ok=True)
    descriptor, name = tempfile.mkstemp(
        prefix=f".{target.stem}.tianyuan-",
        suffix=target.suffix or ".xlsx",
        dir=target.parent,
    )
    os.close(descriptor)
    path = Path(name)
    path.unlink(missing_ok=True)
    return path


def verify_workbook_archive(path: str | Path) -> None:
    path = _ensure_existing_workbook(path)
    try:
        with zipfile.ZipFile(path, "r") as workbook:
            bad_member = workbook.testzip()
    except (OSError, zipfile.BadZipFile) as exc:
        raise ValueError(f"WORKBOOK_INVALID_ZIP: {exc}") from exc
    if bad_member:
        raise ValueError(f"WORKBOOK_CORRUPT_ZIP_MEMBER: {bad_member}")


def _atomic_replace_from(source: Path, target: Path) -> None:
    verify_workbook_archive(source)
    target.parent.mkdir(parents=True, exist_ok=True)
    os.replace(source, target)


def _atomic_copy(source: Path, target: Path) -> None:
    source = _ensure_existing_workbook(source)
    target = _ensure_workbook_path(target)
    if source == target:
        raise ValueError("WORKBOOK_SOURCE_EQUALS_TARGET")
    temporary = _temporary_path(target)
    try:
        shutil.copy2(source, temporary)
        verify_workbook_archive(temporary)
        _atomic_replace_from(temporary, target)
    finally:
        temporary.unlink(missing_ok=True)


def _atomic_transform(path: Path, transform: Callable[[Path], None]) -> None:
    path = _ensure_existing_workbook(path)
    temporary = _temporary_path(path)
    try:
        shutil.copy2(path, temporary)
        transform(temporary)
        verify_workbook_archive(temporary)
        _atomic_replace_from(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def _serialise(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def _serialise_date(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return _serialise(value)


def _header_map(ws, required_headers: list[str]) -> tuple[dict[str, int], list[str]]:
    columns: dict[str, int] = {}
    duplicates: set[str] = set()
    for column, cell in enumerate(ws[3], start=1):
        value = str(cell.value).strip() if cell.value is not None else ""
        if value not in required_headers:
            continue
        if value in columns:
            duplicates.add(value)
        else:
            columns[value] = column
    errors = [f"{ws.title}: 缺少必填表头 `{field}`" for field in required_headers if field not in columns]
    errors.extend(
        f"{ws.title}: 表头 `{field}` 重复，无法判断应读取哪一列"
        for field in sorted(duplicates)
    )
    for field in duplicates:
        columns.pop(field, None)
    return columns, errors


def _all_header_map(ws) -> tuple[dict[str, int], set[str]]:
    columns: dict[str, int] = {}
    duplicates: set[str] = set()
    for column, cell in enumerate(ws[3], start=1):
        value = str(cell.value).strip() if cell.value is not None else ""
        if not value:
            continue
        if value in columns:
            duplicates.add(value)
        else:
            columns[value] = column
    return columns, duplicates


def _template_contract(path: Path) -> dict:
    errors: list[str] = []
    sheets: list[str] = []
    try:
        workbook = load_workbook(path, data_only=False, read_only=True)
    except KeyError as exc:
        raise ValueError(f"TEMPLATE_INCOMPATIBLE: {exc}") from exc
    try:
        sheets = list(workbook.sheetnames)
        for sheet_name in (PARAM_SHEET, STOCK_SHEET, ADDED_SHEET, SUMMARY_SHEET, CHECKS_SHEET):
            if sheet_name not in workbook.sheetnames:
                errors.append(f"缺少工作表 `{sheet_name}`")
        for sheet_name, headers in REQUIRED_HEADERS.items():
            if sheet_name in workbook.sheetnames:
                _, header_errors = _header_map(workbook[sheet_name], headers)
                errors.extend(header_errors)
    finally:
        workbook.close()
    return {"compatible": not errors, "errors": errors, "sheets": sheets}


def _input_snapshot(path: Path) -> dict:
    """Read workbook metadata without changing the workbook or running the model."""
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        parameters = {}
        if PARAM_SHEET in workbook.sheetnames:
            parameter_sheet = workbook[PARAM_SHEET]
            parameter_values = {
                label: parameter_sheet.cell(row, 2).value
                for label, row in PARAMETER_ROWS.items()
            }
            parameters = {
                "valuationDate": _serialise_date(parameter_values.get("评估基准日")),
                "endDate": _serialise_date(parameter_values.get("结束日期")),
                "discountRate": _serialise(parameter_values.get("折现率")),
                "minimumRemainingYears": _serialise(parameter_values.get("最低尚可使用年限")),
                **{label: _serialise_date(value) for label, value in parameter_values.items()},
            }

        counts = {}
        for sheet_name, required_headers in REQUIRED_HEADERS.items():
            if sheet_name not in workbook.sheetnames:
                counts["stock" if sheet_name == STOCK_SHEET else "added"] = 0
                continue
            sheet = workbook[sheet_name]
            columns, _ = _header_map(sheet, required_headers)
            count = 0
            if columns:
                for row in range(4, sheet.max_row + 1):
                    if any(sheet.cell(row, column).value not in (None, "") for column in columns.values()):
                        count += 1
            counts["stock" if sheet_name == STOCK_SHEET else "added"] = count
        return {"parameters": parameters, "counts": counts}
    finally:
        workbook.close()


def preflight_workbook(path: str | Path) -> dict:
    path = _ensure_existing_workbook(path)
    verify_workbook_archive(path)
    lock_path = path.parent / f"~${path.name}"
    contract = _template_contract(path)
    errors = list(contract["errors"])
    engine_errors: list[str] = []
    if contract["compatible"]:
        try:
            _, _, _, engine_errors = read_input(path)
        except (KeyError, ValueError, OSError) as exc:
            engine_errors = [f"输入校验失败：{exc}"]
    errors.extend(engine_errors)
    return {
        "status": "ok" if not errors and not lock_path.exists() else "error",
        "workbook_path": str(path),
        "compatible": contract["compatible"] and not engine_errors,
        "locked": lock_path.exists(),
        "errors": (["工作簿当前被 Office/WPS 锁定，请关闭后重试"] if lock_path.exists() else []) + errors,
        "sheets": contract["sheets"],
    }


def workbook_status(path: str | Path) -> dict:
    path = _ensure_workbook_path(path)
    if not path.exists():
        return {
            "status": "missing",
            "workbook_path": str(path),
            "exists": False,
            "has_results": False,
            "parameters": {},
            "counts": {"stock": 0, "added": 0},
        }
    preflight = preflight_workbook(path)
    snapshot = _input_snapshot(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        result_sheets = [name for name in workbook.sheetnames if name in RESULT_SHEETS]
        summary_ready = SUMMARY_SHEET in workbook.sheetnames and workbook[SUMMARY_SHEET].max_row > 3
        checks_ready = CHECKS_SHEET in workbook.sheetnames and workbook[CHECKS_SHEET].max_row > 3
    finally:
        workbook.close()
    stat = path.stat()
    return {
        "status": "ready" if preflight["status"] == "ok" else "error",
        "workbook_path": str(path),
        "exists": True,
        "size": stat.st_size,
        "modified_ns": stat.st_mtime_ns,
        "has_results": summary_ready and checks_ready,
        "result_sheets": result_sheets,
        **snapshot,
        "preflight": preflight,
    }


def _coerce_input_value(header: str, value: Any) -> Any:
    if isinstance(value, str):
        if "\x00" in value or len(value) > 10000:
            raise ValueError("INPUT_VALUE_INVALID")
        if header in {"启用时间", "预计投入使用时间"}:
            try:
                return date.fromisoformat(value.strip())
            except ValueError:
                return value
    if isinstance(value, (dict, list)):
        raise ValueError("INPUT_VALUE_MUST_BE_SCALAR")
    return value


def _first_blank_row(ws, columns: dict[str, int], start: int = 4) -> int:
    row = start
    while row <= max(ws.max_row, start):
        if not any(ws.cell(row, column).value not in (None, "") for column in columns.values()):
            return row
        row += 1
    return max(ws.max_row + 1, start)


def write_inputs(
    path: str | Path,
    section: str,
    values: dict | list,
    *,
    row_index: int | None = None,
) -> dict:
    path = _ensure_existing_workbook(path)
    section = {"parameters": "params", "parameter": "params", "new": "added"}.get(section, section)
    if section not in {"params", "stock", "added"}:
        raise ValueError("INPUT_SECTION_NOT_ALLOWED")

    changed_rows: list[int] = []

    def transform(temporary: Path) -> None:
        workbook = load_workbook(temporary, data_only=False)
        try:
            if section == "params":
                if not isinstance(values, dict):
                    raise ValueError("PARAMETERS_MUST_BE_OBJECT")
                ws = workbook[PARAM_SHEET]
                for label, value in values.items():
                    if label not in PARAMETER_ROWS:
                        raise ValueError(f"PARAMETER_NOT_ALLOWED: {label}")
                    ws.cell(PARAMETER_ROWS[label], 2, _coerce_input_value(label, value))
                changed_rows.append(0)
                workbook.save(temporary)
                return

            sheet_name = STOCK_SHEET if section == "stock" else ADDED_SHEET
            ws = workbook[sheet_name]
            required = REQUIRED_HEADERS[sheet_name]
            columns, errors = _header_map(ws, required)
            if errors:
                raise ValueError("TEMPLATE_INCOMPATIBLE: " + "；".join(errors))
            all_columns, metadata_duplicates = _all_header_map(ws)
            if metadata_duplicates:
                raise ValueError(
                    "TEMPLATE_INCOMPATIBLE: 表头重复，无法安全写入："
                    + "、".join(sorted(metadata_duplicates))
                )
            rows = values if isinstance(values, list) else values.get("rows") if isinstance(values, dict) and isinstance(values.get("rows"), list) else [values]
            if not rows or any(not isinstance(item, dict) for item in rows):
                raise ValueError("INPUT_ROWS_MUST_BE_OBJECTS")
            for offset, item in enumerate(rows):
                target_row = item.get("rowIndex") if isinstance(item.get("rowIndex"), int) else None
                if target_row is None and row_index is not None:
                    target_row = row_index + offset
                if target_row is None:
                    target_row = _first_blank_row(ws, columns)
                if target_row < 4 or target_row > 100000:
                    raise ValueError("INPUT_ROW_INDEX_INVALID")
                for header, value in item.items():
                    if header in {"rowIndex", "rows"}:
                        continue
                    if header not in all_columns:
                        raise ValueError(f"INPUT_FIELD_NOT_ALLOWED: {header}")
                    ws.cell(target_row, all_columns[header], _coerce_input_value(header, value))
                changed_rows.append(target_row)
            workbook.save(temporary)
        finally:
            workbook.close()

    _atomic_transform(path, transform)
    return {
        "status": "written",
        "workbook_path": str(path),
        "section": section,
        "rows": changed_rows,
        "preflight": preflight_workbook(path),
    }


def _read_value(value: Any) -> Any:
    return _serialise(value)


def read_table_page(
    path: str | Path,
    sheet_name: str,
    *,
    page: int = 1,
    page_size: int = 50,
    column_page: int = 1,
    column_page_size: int = 40,
    column_start: int | None = None,
    column_count: int | None = None,
    data_only: bool = True,
) -> dict:
    path = _ensure_existing_workbook(path)
    verify_workbook_archive(path)
    page = max(1, min(int(page), 100000))
    page_size = max(1, min(int(page_size), 500))
    column_page = max(1, min(int(column_page), 100000))
    column_page_size = max(1, min(int(column_page_size), 200))
    workbook = load_workbook(path, read_only=True, data_only=data_only)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"RESULT_SHEET_NOT_FOUND: {sheet_name}")
        ws = workbook[sheet_name]
        if sheet_name == MONTHLY_DETAIL_SHEET:
            data_start = 6
        elif sheet_name == DETAIL_PROCESS_SHEET:
            # The detail-process sheet contains an asset metadata block before
            # its actual monthly table. Return the table segment so callers can
            # render rows and headers without loading the whole 200-year sheet.
            data_start = None
            for candidate_row in range(1, min(ws.max_row, 500) + 1):
                values = [ws.cell(candidate_row, column).value for column in range(1, min(ws.max_column, 20) + 1)]
                if "日期列表" in values and "折现期（月）" in values:
                    data_start = candidate_row + 1
                    break
            data_start = data_start or 1
        else:
            data_start = 4
        total_rows = max(0, ws.max_row - data_start + 1)
        total_columns = max(0, ws.max_column)
        row_start = data_start + (page - 1) * page_size
        row_end = min(ws.max_row, row_start + page_size - 1)
        if column_start is not None:
            column_start = max(1, min(int(column_start), total_columns or 1))
            effective_column_count = max(1, min(int(column_count or column_page_size), 200))
            column_end = min(total_columns, column_start + effective_column_count - 1)
        else:
            column_start = 1 + (column_page - 1) * column_page_size
            column_end = min(total_columns, column_start + column_page_size - 1)
        if row_start > ws.max_row or column_start > total_columns:
            rows: list[list[Any]] = []
        else:
            rows = [
                [_read_value(value) for value in row]
                for row in ws.iter_rows(
                    min_row=row_start,
                    max_row=row_end,
                    min_col=column_start,
                    max_col=column_end,
                    values_only=True,
                )
            ]
        headers = []
        header_row = 3
        if sheet_name == DETAIL_PROCESS_SHEET and data_start > 1:
            header_row = data_start - 1
        if column_start <= total_columns:
            headers = [
                _read_value(value)
                for row in ws.iter_rows(
                    min_row=header_row,
                    max_row=header_row,
                    min_col=column_start,
                    max_col=column_end,
                    values_only=True,
                )
                for value in row
            ]
        return {
            "status": "ok",
            "workbook_path": str(path),
            "sheet": sheet_name,
            "page": page,
            "page_size": page_size,
            "column_page": column_page,
            "column_page_size": column_page_size,
            "requested_column_start": column_start,
            "requested_column_count": column_count,
            "data_start_row": data_start,
            "row_start": row_start if rows else None,
            "row_end": row_end if rows else None,
            "column_start": column_start if headers else None,
            "column_end": column_end if headers else None,
            "total_rows": total_rows,
            "total_columns": total_columns,
            "headers": headers,
            "rows": rows,
            "has_next_page": row_end < ws.max_row if rows else False,
            "has_next_column_page": column_end < total_columns if headers else False,
        }
    finally:
        workbook.close()


def export_workbook(template_path: str | Path, output_path: str | Path, *, overwrite: bool = False) -> dict:
    source = _ensure_existing_workbook(template_path)
    target = _ensure_workbook_path(output_path)
    if source == target:
        raise ValueError("EXPORT_TARGET_SAME_AS_SOURCE")
    if target.exists() and not overwrite:
        raise FileExistsError(f"EXPORT_TARGET_EXISTS: {target}")
    _atomic_copy(source, target)
    digest = hashlib.sha256(target.read_bytes()).hexdigest()
    workbook = load_workbook(target, read_only=True, data_only=True)
    try:
        sheets = list(workbook.sheetnames)
    finally:
        workbook.close()
    return {
        "status": "exported",
        "workbook_path": str(source),
        "output_path": str(target),
        "size": target.stat().st_size,
        "sha256": digest,
        "readback": {
            "preflight": preflight_workbook(target),
            "sheets": sheets,
            "summary": read_summary(target),
        },
    }


def prepare_workbook(reset: bool = False, output_path: str | Path | None = None) -> dict:
    target = _ensure_workbook_path(output_path or WORKBOOK_PATH)
    target.parent.mkdir(parents=True, exist_ok=True)
    if not target.exists():
        _atomic_copy(CANONICAL_TEMPLATE, target)
        source = "skill-asset"
    elif reset:
        _atomic_copy(CANONICAL_TEMPLATE, target)
        source = "skill-asset-reset"
    else:
        source = "existing-working-workbook"
    return {
        "status": "ready_for_input",
        "skill_root": str(SKILL_ROOT),
        "workbook_path": str(target),
        "source": source,
        "next_step": "请填写该工作簿，填写完成后关闭文件，再让 AI 运行计算。",
    }


def read_summary(workbook_path: Path) -> dict:
    wb = load_workbook(workbook_path, data_only=True, read_only=True)
    try:
        summary_ws = wb["预测汇总"]
        checks_ws = wb["检查结果"]

        header = [cell for cell in next(summary_ws.iter_rows(min_row=3, max_row=3, values_only=True))]
        years = [value for value in header[1:] if isinstance(value, int)]

        summary_headers = [value for value in header if value not in (None, "")]
        summary_rows = []
        expense_headers = []
        expense_rows = []
        in_expense_section = False

        metrics: dict[str, dict[str, float]] = {}
        for row in summary_ws.iter_rows(min_row=4, values_only=True):
            if not row or row[0] in (None, "费用科目"):
                if row and row[0] == "费用科目":
                    in_expense_section = True
                    expense_headers = [value for value in row if value not in (None, "")]
                continue
            label = row[0]
            if in_expense_section:
                values = [row[index] if index < len(row) else None for index in range(len(expense_headers))]
                expense_rows.append(values)
                continue
            if label in {"模型状态", "合计", "主营业务成本", "其他业务成本", "制造费用", "研发费用", "管理费用", "销售费用"}:
                continue
            summary_rows.append([row[index] if index < len(summary_headers) else None for index in range(len(summary_headers))])
            values = {}
            for idx, year in enumerate(years, start=1):
                values[str(year)] = row[idx]
            values["永续期"] = row[len(years) + 1]
            metrics[label] = values

        checks = []
        for row in checks_ws.iter_rows(min_row=4, values_only=True):
            if not row or row[0] is None:
                continue
            checks.append(
                {
                    "检查项": row[0],
                    "实际值": row[1],
                    "期望值": row[2],
                    "状态": row[3],
                    "说明": row[4],
                }
            )

        validation = next((item for item in checks if item["检查项"] == "输入校验"), None)
        checks_headers = ["检查项", "实际值", "期望值", "状态", "说明"]
        checks_rows = [[item.get(header) for header in checks_headers] for item in checks]
        return {
            "years": years,
            "metrics": metrics,
            "expense_summary": {
                "headers": expense_headers or summary_headers,
                "rows": expense_rows,
            },
            "tables": {
                "预测汇总": {"headers": summary_headers, "rows": summary_rows},
                "费用科目汇总": {"headers": expense_headers or summary_headers, "rows": expense_rows},
                "检查结果": {"headers": checks_headers, "rows": checks_rows},
            },
            "checks": checks,
            "input_validation": validation,
        }
    finally:
        wb.close()


def run_workflow(
    *,
    workbook_path: Path | None = None,
    output_path: Path | None = None,
    include_details: bool = False,
    detail_asset_id: str | None = None,
    detail_asset_kind: str | None = None,
) -> dict:
    if workbook_path is None:
        prepare_info = prepare_workbook(reset=False)
        workbook_path = WORKBOOK_PATH
    else:
        workbook_path = workbook_path.expanduser().resolve()
        if not workbook_path.exists():
            raise FileNotFoundError(f"未找到指定输入工作簿：{workbook_path}")
        prepare_info = {"status": "provided_input", "workbook_path": str(workbook_path)}
    output_path = _ensure_workbook_path(output_path or workbook_path)
    workbook_path = _ensure_existing_workbook(workbook_path)
    temporary_output = _temporary_path(output_path)
    cmd = [
        sys.executable,
        "-m",
        "depreciation_forecast",
        "run",
        "--input",
        str(workbook_path),
        "--output",
        str(temporary_output),
    ]
    if include_details:
        cmd.append("--include-details")
    if detail_asset_id:
        cmd.extend(["--detail-asset-id", detail_asset_id])
    if detail_asset_kind:
        cmd.extend(["--detail-asset-kind", detail_asset_kind])
    try:
        result = subprocess.run(
            cmd,
            cwd=SKILL_ROOT,
            env=_python_env(),
            check=True,
            capture_output=True,
            text=True,
        )
        verify_workbook_archive(temporary_output)
        _atomic_replace_from(temporary_output, output_path)
    except subprocess.CalledProcessError as exc:
        return {
            "status": "error",
            "skill_root": str(SKILL_ROOT),
            "workbook_path": str(workbook_path),
            "stdout": exc.stdout,
            "stderr": exc.stderr,
            "message": "运行失败。若工作簿正在打开，请先关闭文件后再重试。",
        }
    except (OSError, ValueError) as exc:
        return {
            "status": "error",
            "skill_root": str(SKILL_ROOT),
            "workbook_path": str(workbook_path),
            "message": str(exc),
        }
    finally:
        temporary_output.unlink(missing_ok=True)

    summary = read_summary(output_path)
    return {
        "status": "ok",
        "skill_root": str(SKILL_ROOT),
        "workbook_path": str(output_path),
        "include_details": include_details,
        "detail_asset_id": detail_asset_id,
        "detail_asset_kind": detail_asset_kind,
        "stdout": result.stdout.strip(),
        "prepare": prepare_info,
        **summary,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="折与资 skill workflow helper")
    subparsers = parser.add_subparsers(dest="command", required=True)

    prepare = subparsers.add_parser("prepare", help="准备工作簿模板")
    prepare.add_argument("--reset", action="store_true", help="重置工作簿为模板状态")
    prepare.add_argument("--output", help="工作簿路径；留空时使用默认 Downloads 路径")

    status = subparsers.add_parser("status", help="读取工作簿当前状态")
    status.add_argument("--input", required=True, help="工作簿路径")

    preflight = subparsers.add_parser("preflight", help="校验模板兼容性和输入字段")
    preflight.add_argument("--input", required=True, help="工作簿路径")

    write = subparsers.add_parser("write", help="安全写入输入字段")
    write.add_argument("--input", required=True, help="工作簿路径")
    write.add_argument("--section", choices=["params", "parameters", "stock", "added", "new"], required=True)
    write.add_argument("--values", required=True, help="JSON 对象或对象数组")
    write.add_argument("--row-index", type=int, help="存量/新增输入的起始行号")

    read = subparsers.add_parser("read", help="分页读取结果工作表")
    read.add_argument("--input", required=True, help="工作簿路径")
    read.add_argument("--sheet", choices=["summary", "annual", "monthly", "detail-process"], required=True)
    read.add_argument("--page", type=int, default=1)
    read.add_argument("--page-size", type=int, default=50)
    read.add_argument("--column-page", type=int, default=1)
    read.add_argument("--column-page-size", type=int, default=40)
    read.add_argument("--column-start", type=int, help="直接指定 1-based 时间列起点")
    read.add_argument("--column-count", type=int, help="直接指定读取的时间列数量")
    read.add_argument("--formulas", action="store_true", help="读取公式而非缓存值")

    export = subparsers.add_parser("export", help="原子导出并回读底稿")
    export.add_argument("--input", required=True, help="结果工作簿路径")
    export.add_argument("--output", required=True, help="导出目标路径")
    export.add_argument("--overwrite", action="store_true", help="明确允许覆盖已存在目标")

    run = subparsers.add_parser("run", help="运行预测并输出摘要 JSON")
    run.add_argument("--input", help="已填写的输入工作簿路径；不填时使用默认 Downloads 工作簿")
    run.add_argument("--output", help="输出路径；不填时直接回写输入工作簿")
    run.add_argument("--with-details", action="store_true", help="验证模式：同时生成年度/月度明细")
    run.add_argument("--detail-asset-id", help="输出指定资产的详细计算过程")
    run.add_argument("--detail-asset-kind", choices=["存量", "新增"], help="与 detail-asset-id 一起使用，限定资产类别")

    args = parser.parse_args()
    if args.command == "prepare":
        payload = prepare_workbook(reset=args.reset, output_path=args.output)
    elif args.command == "status":
        payload = workbook_status(args.input)
    elif args.command == "preflight":
        payload = preflight_workbook(args.input)
    elif args.command == "write":
        try:
            values = json.loads(args.values)
        except json.JSONDecodeError as exc:
            raise SystemExit(f"INPUT_JSON_INVALID: {exc}") from exc
        payload = write_inputs(args.input, args.section, values, row_index=args.row_index)
    elif args.command == "read":
        sheet_map = {
            "summary": SUMMARY_SHEET,
            "annual": ANNUAL_DETAIL_SHEET,
            "monthly": MONTHLY_DETAIL_SHEET,
            "detail-process": DETAIL_PROCESS_SHEET,
        }
        if args.sheet == "summary":
            payload = {"status": "ok", "workbook_path": str(_ensure_existing_workbook(args.input)), "summary": read_summary(Path(args.input))}
        else:
            payload = read_table_page(
                args.input,
                sheet_map[args.sheet],
                page=args.page,
                page_size=args.page_size,
                column_page=args.column_page,
                column_page_size=args.column_page_size,
                column_start=args.column_start,
                column_count=args.column_count,
                data_only=not args.formulas,
            )
    elif args.command == "export":
        payload = export_workbook(args.input, args.output, overwrite=args.overwrite)
    else:
        payload = run_workflow(
            workbook_path=Path(args.input) if args.input else None,
            output_path=Path(args.output) if args.output else None,
            include_details=args.with_details,
            detail_asset_id=args.detail_asset_id,
            detail_asset_kind=args.detail_asset_kind,
        )
    print(json.dumps(payload, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except SystemExit:
        raise
    except Exception as exc:
        print(json.dumps({"status": "error", "reason": str(exc)}, ensure_ascii=False, indent=2))
        raise SystemExit(1) from exc
