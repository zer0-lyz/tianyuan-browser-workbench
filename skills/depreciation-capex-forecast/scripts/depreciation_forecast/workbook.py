from __future__ import annotations

import calendar
import random
import re
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils.datetime import from_excel
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .model import (
    METRICS,
    AddedAsset,
    AssetForecast,
    ForecastParams,
    ForecastResult,
    StockAsset,
    as_date,
    completed_months,
    eomonth,
    forecast_portfolio,
    forecast_added,
    forecast_stock,
    money,
)


PARAM_SHEET = "参数设定"
STOCK_SHEET = "存量资产输入"
ADDED_SHEET = "新增资产输入"
SUMMARY_SHEET = "预测汇总"
EXPENSE_SHEET = "费用科目汇总"
DETAIL_SHEET = "资产结果明细（年度）"
LEGACY_DETAIL_SHEET = "资产结果明细"
CHECKS_SHEET = "检查结果"
MONTHLY_DETAIL_SHEET = "资产结果明细（月度）"
DETAIL_PROCESS_SHEET = "详细过程"
DETAIL_HORIZON_YEARS = 200

STOCK_HEADERS = [
    "序号",
    "情形描述",
    "公司主体",
    "资产科目",
    "名称",
    "账面原值",
    "账面净值",
    "评估原值",
    "启用时间",
    "折旧年限",
    "经济耐用年限",
    "预计尚可使用年限",
    "更新后折旧年限",
    "更新后经济耐用年限",
    "残值率",
    "费用科目",
    "折旧摊销",
    "进项税率",
    "是否需要更新",
]

ADDED_HEADERS = [
    "序号",
    "情形描述",
    "公司主体",
    "资产科目",
    "名称",
    "在建工程账面价值",
    "总投资额（不含税）",
    "预计投入使用时间",
    "更新后折旧年限",
    "更新后经济耐用年限",
    "残值率",
    "费用科目",
    "折旧摊销",
    "进项税率",
    "是否需要更新",
]

EXPENSE_ACCOUNTS = ["主营业务成本", "其他业务成本", "销售费用", "管理费用", "研发费用"]
METRIC_ATTRS = {
    "折旧摊销": "depreciation",
    "追加资本性支出": "added_capex",
    "更新资本性支出": "renewal_capex",
    "更新支出进项税": "input_tax",
    "残值回收": "salvage_recovery",
    "残值回收销项税": "output_tax",
}

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
SECTION_FILL = PatternFill("solid", fgColor="D9EAF7")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
OK_FILL = PatternFill("solid", fgColor="C6EFCE")
ERROR_FILL = PatternFill("solid", fgColor="FFC7CE")
THIN_BLUE = Side(style="thin", color="9EBDD7")
MEDIUM_BLUE = Side(style="medium", color="1F4E78")
TABLE_BORDER = Border(bottom=THIN_BLUE)
TOTAL_BORDER = Border(top=MEDIUM_BLUE)
INPUT_FONT = Font(color="0000FF")
FORMULA_FONT = Font(color="000000")


def _style_header(ws, row: int, start: int, end: int) -> None:
    for col in range(start, end + 1):
        cell = ws.cell(row, col)
        cell.fill = HEADER_FILL
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _auto_width(ws, min_width: int = 10, max_width: int = 28) -> None:
    for col in range(1, ws.max_column + 1):
        values = [str(ws.cell(row, col).value or "") for row in range(1, min(ws.max_row, 120) + 1)]
        width = min(max(max(len(value) for value in values) + 2, min_width), max_width)
        ws.column_dimensions[get_column_letter(col)].width = width


def _sheet_title(ws, title: str, end_col: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
    cell = ws.cell(1, 1, title)
    cell.fill = HEADER_FILL
    cell.font = Font(color="FFFFFF", bold=True, size=14)
    cell.alignment = Alignment(horizontal="left")


def _add_validations(ws, max_row: int, kind_col: str, renew_col: str) -> None:
    kind = DataValidation(type="list", formula1='"折旧,摊销"', allow_blank=False)
    renew = DataValidation(type="list", formula1='"是,否"', allow_blank=False)
    ws.add_data_validation(kind)
    ws.add_data_validation(renew)
    kind.add(f"{kind_col}4:{kind_col}{max_row}")
    renew.add(f"{renew_col}4:{renew_col}{max_row}")


def create_template(path: str | Path, *, include_scenarios: bool = False, prototype_path: str | Path | None = None) -> Path:
    path = Path(path)
    wb = Workbook()
    ws = wb.active
    ws.title = PARAM_SHEET
    _sheet_title(ws, "折旧摊销与资本性支出预测工具 - 公共参数", 4)
    ws.append([])
    ws.append(["参数", "填写值", "说明", "示例"])
    rows = [
        ["评估基准日", date(2024, 5, 31), "预测计算起点", date(2024, 5, 31)],
        ["结束日期", date(2029, 12, 31), "详细预测期终点", date(2029, 12, 31)],
        ["折现率", 0.10, "用于永续期年金化", 0.10],
        ["最低尚可使用年限", 1, "存量资产未填写预计尚可使用年限时使用", 1],
    ]
    for row in rows:
        ws.append(row)
    _style_header(ws, 3, 1, 4)
    for row in range(4, 8):
        ws.cell(row, 2).fill = INPUT_FILL
        ws.cell(row, 2).font = INPUT_FONT
    for cell in ("B4", "B5", "D4", "D5"):
        ws[cell].number_format = "yyyy-mm-dd"
    for cell in ("B6", "D6"):
        ws[cell].number_format = "0.00%"
    ws.freeze_panes = "A4"

    stock = wb.create_sheet(STOCK_SHEET)
    _sheet_title(stock, "存量资产输入", len(STOCK_HEADERS))
    stock.append([])
    stock.append(STOCK_HEADERS)
    _style_header(stock, 3, 1, len(STOCK_HEADERS))
    _add_validations(stock, 10000, "Q", "S")
    stock.freeze_panes = "A4"
    stock.auto_filter.ref = f"A3:S10000"

    added = wb.create_sheet(ADDED_SHEET)
    _sheet_title(added, "新增资产输入", len(ADDED_HEADERS))
    added.append([])
    added.append(ADDED_HEADERS)
    _style_header(added, 3, 1, len(ADDED_HEADERS))
    _add_validations(added, 10000, "M", "O")
    added.freeze_panes = "A4"
    added.auto_filter.ref = f"A3:O10000"

    if include_scenarios:
        if not prototype_path:
            raise ValueError("include_scenarios=True 时必须传入 prototype_path")
        prototype = load_workbook(prototype_path, data_only=True, read_only=True)
        source_stock = prototype["存量资产（情形分析）"]
        source_added = prototype["新增资产（情形分析）"]
        for row in source_stock.iter_rows(min_row=45, max_row=72, min_col=1, max_col=19, values_only=True):
            if row[0] is not None:
                stock.append(list(row))
        for row in source_added.iter_rows(min_row=15, max_row=42, min_col=1, max_col=15, values_only=True):
            if row[0] is not None:
                added.append(list(row))

    for input_ws, headers in ((stock, STOCK_HEADERS), (added, ADDED_HEADERS)):
        for row in input_ws.iter_rows(min_row=4, max_row=max(input_ws.max_row, 200), min_col=1, max_col=len(headers)):
            for cell in row:
                cell.fill = INPUT_FILL
                cell.font = INPUT_FONT
                cell.border = TABLE_BORDER
        _auto_width(input_ws, 10, 22)
        input_ws.column_dimensions["B"].width = 48
    for row in range(4, max(stock.max_row, 200) + 1):
        for col in range(6, 9):
            stock.cell(row, col).number_format = '#,##0.00;[Red](#,##0.00);-'
        stock.cell(row, 9).number_format = "yyyy-mm-dd"
        stock.cell(row, 15).number_format = "0.00%"
        stock.cell(row, 18).number_format = "0.00%"
    for row in range(4, max(added.max_row, 200) + 1):
        for col in range(6, 8):
            added.cell(row, col).number_format = '#,##0.00;[Red](#,##0.00);-'
        added.cell(row, 8).number_format = "yyyy-mm-dd"
        added.cell(row, 11).number_format = "0.00%"
        added.cell(row, 14).number_format = "0.00%"

    for name in (SUMMARY_SHEET, CHECKS_SHEET):
        output = wb.create_sheet(name)
        output["A1"] = "运行脚本后自动生成，请勿手工填写"
        output["A1"].font = Font(italic=True, color="808080")

    _auto_width(ws, 14, 38)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _number(value: Any, field: str, row: int, errors: list[str], *, blank_zero: bool = False) -> float:
    if value is None or value == "":
        if blank_zero:
            return 0.0
        errors.append(f"第 {row} 行 `{field}` 不能为空")
        return 0.0
    try:
        return float(value)
    except (TypeError, ValueError):
        errors.append(f"第 {row} 行 `{field}` 必须为数值，当前值为 {value!r}")
        return 0.0


def _date_value(value: Any, field: str, row: int, errors: list[str]) -> date:
    if isinstance(value, (date, datetime)):
        return as_date(value)
    if isinstance(value, (int, float)):
        try:
            return as_date(from_excel(value))
        except (TypeError, ValueError, OverflowError):
            pass
    if isinstance(value, str):
        try:
            return date.fromisoformat(value.strip())
        except ValueError:
            pass
    errors.append(f"第 {row} 行 `{field}` 必须为日期，当前值为 {value!r}")
    return date(1900, 1, 1)


def _boolean(value: Any, field: str, row: int, errors: list[str]) -> bool:
    if value in ("是", True, 1):
        return True
    if value in ("否", False, 0):
        return False
    errors.append(f"第 {row} 行 `{field}` 必须填写 是 或 否，当前值为 {value!r}")
    return False


def _kind(value: Any, row: int, errors: list[str]) -> str:
    if value in ("折旧", "摊销"):
        return value
    errors.append(f"第 {row} 行 `折旧摊销` 必须填写 折旧 或 摊销，当前值为 {value!r}")
    return "折旧"


def _header_columns(ws, required_headers: list[str], errors: list[str]) -> dict[str, int]:
    """Map required field names from row 3 and tolerate project metadata columns."""
    columns: dict[str, int] = {}
    duplicates: set[str] = set()
    for column, cell in enumerate(ws[3], start=1):
        value = str(cell.value).strip() if cell.value is not None else ""
        if value not in required_headers:
            continue
        if value in columns:
            duplicates.add(value)
        else:
            columns[value] = column
    for field in required_headers:
        if field not in columns:
            errors.append(f"{ws.title}: 缺少必填表头 `{field}`")
    for field in sorted(duplicates):
        errors.append(f"{ws.title}: 表头 `{field}` 重复，无法判断应读取哪一列")
        columns.pop(field, None)
    return columns


def _reference_cell(ws, reference: str):
    return ws[reference.replace("$", "")]


def _formula_number(
    ws,
    cell,
    field: str,
    row: int,
    errors: list[str],
    *,
    valuation_date: date,
) -> float | None:
    formula = str(cell.value).strip()
    direct_reference = re.fullmatch(r"=\s*(\$?[A-Z]{1,3}\$?\d+)\s*", formula, flags=re.IGNORECASE)
    if direct_reference:
        source = _reference_cell(ws, direct_reference.group(1))
        if isinstance(source.value, str) and source.value.startswith("="):
            errors.append(f"第 {row} 行 `{field}` 的公式 `{formula}` 引用了尚未支持的嵌套公式")
            return None
        return _number(source.value, field, row, errors)

    remaining_years = re.fullmatch(
        r"=\s*MAX\s*\(\s*(MIN|MAX)\s*\(\s*(\$?[A-Z]{1,3}\$?\d+)\s*,\s*"
        r"(\$?[A-Z]{1,3}\$?\d+)\s*\)\s*-\s*\(\s*[^-]+-\s*"
        r"(\$?[A-Z]{1,3}\$?\d+)\s*\)\s*/\s*365\s*,\s*0\s*\)\s*",
        formula,
        flags=re.IGNORECASE,
    )
    if remaining_years:
        operator, first_ref, second_ref, service_ref = remaining_years.groups()
        first = _number(_reference_cell(ws, first_ref).value, field, row, errors)
        second = _number(_reference_cell(ws, second_ref).value, field, row, errors)
        service_date = _date_value(_reference_cell(ws, service_ref).value, "启用时间", row, errors)
        base_years = min(first, second) if operator.upper() == "MIN" else max(first, second)
        return max(base_years - (valuation_date - service_date).days / 365, 0.0)

    errors.append(
        f"第 {row} 行 `{field}` 的公式 `{formula}` 未被支持；请改为数值，或使用模板的直接引用/预计尚可使用年限公式"
    )
    return None


def _number_from_cell(
    ws,
    row: int,
    column: int,
    field: str,
    validation_row: int,
    errors: list[str],
    *,
    valuation_date: date,
    blank_zero: bool = False,
) -> float:
    cell = ws.cell(row, column)
    if isinstance(cell.value, str) and cell.value.startswith("="):
        value = _formula_number(ws, cell, field, validation_row, errors, valuation_date=valuation_date)
        return 0.0 if value is None else value
    return _number(cell.value, field, validation_row, errors, blank_zero=blank_zero)


def _row_values(ws, row: int, columns: dict[str, int], headers: list[str]) -> dict[str, Any]:
    return {header: ws.cell(row, columns[header]).value for header in headers if header in columns}


def read_input(path: str | Path) -> tuple[ForecastParams, list[StockAsset], list[AddedAsset], list[str]]:
    # Formula caches are frequently absent after programmatic saves. Read the
    # formulas themselves and evaluate only the template patterns we support.
    wb = load_workbook(path, data_only=False)
    errors: list[str] = []
    param_ws = wb[PARAM_SHEET]
    params = ForecastParams(
        valuation_date=_date_value(param_ws["B4"].value, "评估基准日", 4, errors),
        end_date=_date_value(param_ws["B5"].value, "结束日期", 5, errors),
        discount_rate=_number(param_ws["B6"].value, "折现率", 6, errors),
        minimum_remaining_years=_number(param_ws["B7"].value, "最低尚可使用年限", 7, errors),
    )
    if params.end_date <= params.valuation_date:
        errors.append("`结束日期` 必须晚于 `评估基准日`")
    if params.discount_rate < 0:
        errors.append("`折现率` 不得小于 0")

    stock_assets: list[StockAsset] = []
    stock_ws = wb[STOCK_SHEET]
    stock_columns = _header_columns(stock_ws, STOCK_HEADERS, errors)
    stock_rows = range(4, stock_ws.max_row + 1) if len(stock_columns) == len(STOCK_HEADERS) else ()
    for row in stock_rows:
        values = _row_values(stock_ws, row, stock_columns, STOCK_HEADERS)
        if not any(value not in (None, "") for value in values.values()):
            continue
        row_errors: list[str] = []
        salvage = _number_from_cell(stock_ws, row, stock_columns["残值率"], "残值率", row, row_errors, valuation_date=params.valuation_date, blank_zero=True)
        asset = StockAsset(
            asset_id=str(values["序号"] or f"存量-{row - 3}"),
            scenario=str(values["情形描述"] or ""),
            company=str(values["公司主体"] or ""),
            account=str(values["资产科目"] or ""),
            name=str(values["名称"] or ""),
            original_cost=_number_from_cell(stock_ws, row, stock_columns["账面原值"], "账面原值", row, row_errors, valuation_date=params.valuation_date),
            net_book_value=_number_from_cell(stock_ws, row, stock_columns["账面净值"], "账面净值", row, row_errors, valuation_date=params.valuation_date),
            appraised_original_value=_number_from_cell(stock_ws, row, stock_columns["评估原值"], "评估原值", row, row_errors, valuation_date=params.valuation_date),
            in_service_date=_date_value(values["启用时间"], "启用时间", row, row_errors),
            depreciation_years=_number_from_cell(stock_ws, row, stock_columns["折旧年限"], "折旧年限", row, row_errors, valuation_date=params.valuation_date),
            economic_years=_number_from_cell(stock_ws, row, stock_columns["经济耐用年限"], "经济耐用年限", row, row_errors, valuation_date=params.valuation_date),
            expected_remaining_years=_number_from_cell(stock_ws, row, stock_columns["预计尚可使用年限"], "预计尚可使用年限", row, row_errors, valuation_date=params.valuation_date, blank_zero=True),
            renewed_depreciation_years=_number_from_cell(stock_ws, row, stock_columns["更新后折旧年限"], "更新后折旧年限", row, row_errors, valuation_date=params.valuation_date),
            renewed_economic_years=_number_from_cell(stock_ws, row, stock_columns["更新后经济耐用年限"], "更新后经济耐用年限", row, row_errors, valuation_date=params.valuation_date),
            salvage_rate=salvage,
            expense_account=str(values["费用科目"] or ""),
            depreciation_kind=_kind(values["折旧摊销"], row, row_errors),  # type: ignore[arg-type]
            input_tax_rate=_number_from_cell(stock_ws, row, stock_columns["进项税率"], "进项税率", row, row_errors, valuation_date=params.valuation_date, blank_zero=True),
            renew=_boolean(values["是否需要更新"], "是否需要更新", row, row_errors),
        )
        if asset.original_cost < asset.net_book_value:
            row_errors.append(f"第 {row} 行 `账面原值` 不得小于 `账面净值`")
        if min(asset.original_cost, asset.net_book_value, asset.appraised_original_value) < 0:
            row_errors.append(f"第 {row} 行金额字段不得小于 0")
        if min(
            asset.depreciation_years,
            asset.economic_years,
            asset.renewed_depreciation_years,
            asset.renewed_economic_years,
        ) <= 0:
            row_errors.append(f"第 {row} 行折旧年限和经济耐用年限必须大于 0")
        if asset.expected_remaining_years < 0:
            row_errors.append(f"第 {row} 行 `预计尚可使用年限` 不得小于 0")
        if asset.depreciation_kind == "摊销" and values["残值率"] in (None, ""):
            asset.salvage_rate = 0.0
        _validate_common(asset.salvage_rate, asset.input_tax_rate, row, row_errors)
        if row_errors:
            errors.extend(f"{STOCK_SHEET}: {message}" for message in row_errors)
        else:
            stock_assets.append(asset)

    added_assets: list[AddedAsset] = []
    added_ws = wb[ADDED_SHEET]
    added_columns = _header_columns(added_ws, ADDED_HEADERS, errors)
    added_rows = range(4, added_ws.max_row + 1) if len(added_columns) == len(ADDED_HEADERS) else ()
    for row in added_rows:
        values = _row_values(added_ws, row, added_columns, ADDED_HEADERS)
        if not any(value not in (None, "") for value in values.values()):
            continue
        row_errors = []
        salvage = _number_from_cell(added_ws, row, added_columns["残值率"], "残值率", row, row_errors, valuation_date=params.valuation_date, blank_zero=True)
        asset = AddedAsset(
            asset_id=str(values["序号"] or f"新增-{row - 3}"),
            scenario=str(values["情形描述"] or ""),
            company=str(values["公司主体"] or ""),
            account=str(values["资产科目"] or ""),
            name=str(values["名称"] or ""),
            construction_book_value=_number_from_cell(added_ws, row, added_columns["在建工程账面价值"], "在建工程账面价值", row, row_errors, valuation_date=params.valuation_date, blank_zero=True),
            total_investment=_number_from_cell(added_ws, row, added_columns["总投资额（不含税）"], "总投资额（不含税）", row, row_errors, valuation_date=params.valuation_date),
            in_service_date=_date_value(values["预计投入使用时间"], "预计投入使用时间", row, row_errors),
            renewed_depreciation_years=_number_from_cell(added_ws, row, added_columns["更新后折旧年限"], "更新后折旧年限", row, row_errors, valuation_date=params.valuation_date),
            renewed_economic_years=_number_from_cell(added_ws, row, added_columns["更新后经济耐用年限"], "更新后经济耐用年限", row, row_errors, valuation_date=params.valuation_date),
            salvage_rate=salvage,
            expense_account=str(values["费用科目"] or ""),
            depreciation_kind=_kind(values["折旧摊销"], row, row_errors),  # type: ignore[arg-type]
            input_tax_rate=_number_from_cell(added_ws, row, added_columns["进项税率"], "进项税率", row, row_errors, valuation_date=params.valuation_date, blank_zero=True),
            renew=_boolean(values["是否需要更新"], "是否需要更新", row, row_errors),
        )
        if asset.construction_book_value > asset.total_investment:
            row_errors.append(f"第 {row} 行 `在建工程账面价值` 不得大于 `总投资额（不含税）`")
        if min(asset.construction_book_value, asset.total_investment) < 0:
            row_errors.append(f"第 {row} 行金额字段不得小于 0")
        if min(asset.renewed_depreciation_years, asset.renewed_economic_years) <= 0:
            row_errors.append(f"第 {row} 行折旧年限和经济耐用年限必须大于 0")
        if asset.in_service_date <= params.valuation_date:
            row_errors.append(f"第 {row} 行 `预计投入使用时间` 必须晚于评估基准日")
        if asset.depreciation_kind == "摊销" and values["残值率"] in (None, ""):
            asset.salvage_rate = 0.0
        _validate_common(asset.salvage_rate, asset.input_tax_rate, row, row_errors)
        if row_errors:
            errors.extend(f"{ADDED_SHEET}: {message}" for message in row_errors)
        else:
            added_assets.append(asset)
    return params, stock_assets, added_assets, errors


def _validate_common(salvage_rate: float, input_tax_rate: float, row: int, errors: list[str]) -> None:
    if not 0 <= salvage_rate <= 1:
        errors.append(f"第 {row} 行 `残值率` 必须在 0 到 1 之间")
    if not 0 <= input_tax_rate <= 1:
        errors.append(f"第 {row} 行 `进项税率` 必须在 0 到 1 之间")


def _clear_sheet(ws) -> None:
    ws.delete_rows(1, ws.max_row)


def _write_matrix(ws, rows: Iterable[Iterable[Any]]) -> None:
    for row in rows:
        ws.append(list(row))


def _add_years(value: date, years: int) -> date:
    year = value.year + years
    last_day = calendar.monthrange(year, value.month)[1]
    return date(year, value.month, min(value.day, last_day))


def _format_output(ws) -> None:
    ws.freeze_panes = "A4"
    ws.sheet_view.showGridLines = False
    _auto_width(ws, 12, 24)
    for row in ws.iter_rows(min_row=4):
        for cell in row:
            cell.border = TABLE_BORDER
            if isinstance(cell.value, (int, float)):
                cell.number_format = '#,##0.00;[Red](#,##0.00);-'
    if ws.max_row >= 3:
        _style_header(ws, 3, 1, ws.max_column)
        for cell in ws[3]:
            if isinstance(cell.value, (int, float)):
                cell.number_format = "0"


def _month_timeline(start: date, end: date) -> list[date]:
    values = [start]
    cursor = eomonth(start, 1)
    while cursor <= end:
        values.append(cursor)
        cursor = eomonth(cursor, 1)
    return values


def _forecast_years(valuation_date: date, end_date: date) -> list[int]:
    first_year = eomonth(valuation_date, 1).year
    return list(range(first_year, end_date.year + 1))


def _month_index(start: date, end: date) -> int:
    start = as_date(start)
    end = as_date(end)
    return (end.year - start.year) * 12 + end.month - start.month


def _metric_amount(value: Any, metric: str) -> float:
    return float(getattr(value, METRIC_ATTRS[metric], 0.0))


def _discount_factor_for_month(month: date, cutoff_end_date: date, monthly_rate: float) -> float:
    if month <= cutoff_end_date:
        return 1.0
    discount_months = _month_index(cutoff_end_date, month)
    return 1 / (1 + monthly_rate) ** discount_months if discount_months else 1.0


def _annualize_perpetual(total_present_value: float, annual_rate: float) -> float:
    return money(total_present_value * annual_rate)


def _asset_yearly_totals(
    asset: AssetForecast,
    years: list[int],
    cutoff_end_date: date,
    monthly_rate: float,
    *,
    discounted: bool,
) -> dict[tuple[str, int], float]:
    result: dict[tuple[str, int], float] = defaultdict(float)
    year_set = set(years)
    for value in asset.monthly.values():
        if value.month.year not in year_set:
            continue
        if discounted:
            if value.month <= cutoff_end_date:
                continue
            factor = _discount_factor_for_month(value.month, cutoff_end_date, monthly_rate)
        else:
            factor = 1.0
        for metric in METRIC_ATTRS:
            result[metric, value.month.year] += _metric_amount(value, metric) * factor
    return defaultdict(float, {key: money(amount) for key, amount in result.items()})


def _asset_perpetual_present_values(
    asset: AssetForecast,
    cutoff_end_date: date,
    monthly_rate: float,
    annual_rate: float,
) -> dict[str, float]:
    result: dict[str, float] = defaultdict(float)
    for value in asset.monthly.values():
        if value.month <= cutoff_end_date:
            continue
        factor = _discount_factor_for_month(value.month, cutoff_end_date, monthly_rate)
        for metric in METRIC_ATTRS:
            result[metric] += _metric_amount(value, metric) * factor
    result["资本性支出"] = result["追加资本性支出"] + result["更新资本性支出"]
    return defaultdict(float, {key: _annualize_perpetual(amount, annual_rate) for key, amount in result.items()})


def _output_totals(
    assets: list[AssetForecast],
    params: ForecastParams,
    cutoff_end_date: date,
) -> tuple[dict[tuple[str, int], float], dict[str, float], dict[tuple[str, int | str], float]]:
    annual: dict[tuple[str, int], float] = defaultdict(float)
    perpetual: dict[str, float] = defaultdict(float)
    expense: dict[tuple[str, int | str], float] = defaultdict(float)
    year_set = set(_forecast_years(params.valuation_date, cutoff_end_date))
    for asset in assets:
        for value in asset.monthly.values():
            for metric in METRIC_ATTRS:
                amount = _metric_amount(value, metric)
                if value.month <= cutoff_end_date:
                    if value.month.year in year_set:
                        annual[metric, value.month.year] += amount
                        if metric == "折旧摊销":
                            expense[asset.expense_account, value.month.year] += amount
        for metric in METRIC_ATTRS:
            amount = float(asset.perpetual.get(metric, 0.0))
            perpetual[metric] += amount
            if metric == "折旧摊销":
                expense[asset.expense_account, "永续期"] += amount
    for year in year_set:
        annual["资本性支出", year] = annual["追加资本性支出", year] + annual["更新资本性支出", year]
    perpetual["资本性支出"] = perpetual["追加资本性支出"] + perpetual["更新资本性支出"]
    perpetual = defaultdict(float, {key: money(amount) for key, amount in perpetual.items()})
    expense = defaultdict(
        float,
        {
            key: money(amount)
            for key, amount in expense.items()
        },
    )
    return (
        defaultdict(float, {key: money(amount) for key, amount in annual.items()}),
        perpetual,
        expense,
    )


def _detail_cell_value(asset: AssetForecast, key: str) -> Any:
    return asset.debug.get(key, "")


def _write_detail_process_sheet(
    ws,
    asset: AssetForecast,
    params: ForecastParams,
    *,
    cutoff_end_date: date,
    horizon_years: int,
) -> None:
    horizon_end_date = _add_years(cutoff_end_date, horizon_years)
    timeline = _month_timeline(params.valuation_date, horizon_end_date)
    years = _forecast_years(params.valuation_date, cutoff_end_date)

    ws.delete_rows(1, ws.max_row)
    _sheet_title(ws, f"详细过程 - {asset.asset_kind} {asset.asset_id} {asset.name}", 25)
    ws["A2"] = f"假设均在月底更新；永续期 = 详细预测期后{horizon_years}年折现现值合计 × 折现率"
    ws["A2"].font = Font(italic=True, color="808080")

    summary_start_col = 17  # Q
    summary_headers = ["项目", *years, "永续期", "公式"]
    for offset, value in enumerate(summary_headers, start=summary_start_col):
        cell = ws.cell(3, offset, value)
        cell.fill = HEADER_FILL
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    monthly_headers = [
        "序号",
        "年份",
        "日期列表",
        "追加资本性支出",
        "追加资本性支出（至结束日期现值）",
        "更新资本性支出",
        "更新资本性支出（至结束日期现值）",
        "残值回收",
        "残值回收（至结束日期现值）",
        "更新金额（净额）",
        "更新金额（至结束日期现值）",
        "进项税",
        "进项税（至结束日期现值）",
        "销项税",
        "销项税（至结束日期现值）",
        "月折旧额",
        "折现期（月）",
        "折现系数",
        "更新金额现值",
        "月折旧额现值",
    ]

    info_rows = [
        ("类别", asset.asset_kind),
        ("序号", asset.asset_id),
        ("公司主体", asset.company),
        ("资产科目", asset.account),
        ("名称", asset.name),
        ("费用科目", asset.expense_account),
        ("评估基准日", params.valuation_date),
        ("结束日期", cutoff_end_date),
        ("详细预测期后年数", horizon_years),
        ("详细预测期后截止日", horizon_end_date),
        ("折现率", params.discount_rate),
        ("月折现率", params.monthly_discount_rate),
    ]
    row = 4
    for label, value in info_rows:
        ws.cell(row, 1, label)
        ws.cell(row, 2, value)
        ws.cell(row, 1).fill = SECTION_FILL
        ws.cell(row, 1).font = Font(bold=True)
        ws.cell(row, 2).fill = INPUT_FILL
        row += 1

    value_rows: dict[str, int] = {label: idx for idx, (label, _) in enumerate(info_rows, start=4)}
    derived: list[tuple[str, Any]] = []
    if asset.asset_kind == "存量":
        derived = [
            ("启用时间", _detail_cell_value(asset, "启用时间")),
            ("账面原值", _detail_cell_value(asset, "账面原值")),
            ("账面净值", _detail_cell_value(asset, "账面净值")),
            ("评估原值", _detail_cell_value(asset, "评估原值")),
            ("折旧年限", _detail_cell_value(asset, "折旧年限")),
            ("经济耐用年限", _detail_cell_value(asset, "经济耐用年限")),
            ("预计尚可使用年限", _detail_cell_value(asset, "预计尚可使用年限")),
            ("已使用月份", _detail_cell_value(asset, "已使用月份")),
            ("剩余尚可使用月份", _detail_cell_value(asset, "剩余尚可使用月份")),
            ("预测开始月", ""),
            ("初始折旧月数", _detail_cell_value(asset, "初始折旧月数")),
            ("初始每月折旧", _detail_cell_value(asset, "初始每月折旧")),
            ("首次更新月", ""),
            ("更新后折旧年限", _detail_cell_value(asset, "更新后折旧年限")),
            ("更新后经济耐用年限", _detail_cell_value(asset, "更新后经济耐用年限")),
            ("更新折旧月数", _detail_cell_value(asset, "更新折旧月数")),
            ("更新后经济耐用月数", ""),
            ("更新后每月折旧", _detail_cell_value(asset, "更新后每月折旧")),
            ("原资产残值", _detail_cell_value(asset, "残值")),
            ("更新资产残值", ""),
            ("首次更新净支出", ""),
            ("后续更新净支出", ""),
            ("残值率", _detail_cell_value(asset, "残值率")),
            ("进项税率", _detail_cell_value(asset, "进项税率")),
            ("是否需要更新", "是" if _detail_cell_value(asset, "是否需要更新") else "否"),
            ("折旧摊销方式", _detail_cell_value(asset, "折旧摊销方式")),
            ("永续期折旧", _detail_cell_value(asset, "永续期折旧")),
        ]
    else:
        derived = [
            ("预计投入使用时间", _detail_cell_value(asset, "预计投入使用时间")),
            ("在建工程账面价值", _detail_cell_value(asset, "在建工程账面价值")),
            ("总投资额", _detail_cell_value(asset, "总投资额")),
            ("投入使用月", _detail_cell_value(asset, "投入使用月")),
            ("追加资本性支出", _detail_cell_value(asset, "追加资本性支出")),
            ("折旧月数", _detail_cell_value(asset, "折旧月数")),
            ("每月折旧", _detail_cell_value(asset, "每月折旧")),
            ("经济耐用月数", _detail_cell_value(asset, "经济耐用月数")),
            ("折旧开始月", ""),
            ("首次更新月", ""),
            ("剩余尚可使用月份", _detail_cell_value(asset, "剩余尚可使用月份")),
            ("残值", _detail_cell_value(asset, "残值")),
            ("后续更新净支出", ""),
            ("残值率", _detail_cell_value(asset, "残值率")),
            ("进项税率", _detail_cell_value(asset, "进项税率")),
            ("是否需要更新", "是" if _detail_cell_value(asset, "是否需要更新") else "否"),
            ("折旧摊销方式", _detail_cell_value(asset, "折旧摊销方式")),
            ("永续期折旧", _detail_cell_value(asset, "永续期折旧")),
        ]
    for label, value in derived:
        ws.cell(row, 1, label)
        ws.cell(row, 2, value)
        ws.cell(row, 1).fill = SECTION_FILL
        ws.cell(row, 1).font = Font(bold=True)
        value_rows[label] = row
        row += 1

    def ref(label: str) -> str:
        return f"$B${value_rows[label]}"

    if asset.asset_kind == "存量":
        ws.cell(value_rows["预测开始月"], 2, f'=EOMONTH({ref("评估基准日")},1)')
        ws.cell(value_rows["首次更新月"], 2, f'=EOMONTH({ref("评估基准日")},{ref("剩余尚可使用月份")})')
        ws.cell(value_rows["更新后经济耐用月数"], 2, f'=ROUND({ref("更新后经济耐用年限")}*12,0)')
        ws.cell(value_rows["更新资产残值"], 2, f'=IF({ref("是否需要更新")}="是",{ref("评估原值")}*{ref("残值率")},0)')
        ws.cell(value_rows["首次更新净支出"], 2, f'=IF({ref("是否需要更新")}="是",{ref("评估原值")}-{ref("原资产残值")},0)')
        ws.cell(value_rows["后续更新净支出"], 2, f'=IF({ref("是否需要更新")}="是",{ref("评估原值")}-{ref("更新资产残值")},0)')
        renewal_trigger = (
            f'IF({ref("是否需要更新")}<>"是",0,'
            f'IF($CROW={ref("首次更新月")},1,'
            f'IF($CROW<={ref("首次更新月")},0,'
            f'IF(MOD(DATEDIF({ref("首次更新月")},$CROW,"m"),{ref("更新后经济耐用月数")})=0,2,0))))'
        )
        dep_offset = f'IF({ref("折旧摊销方式")}="折旧",1,0)'
        initial_dep_formula = (
            f'IF(AND($CROW>={ref("预测开始月")},DATEDIF({ref("预测开始月")},$CROW,"m")<{ref("初始折旧月数")}),'
            f'{ref("初始每月折旧")},0)'
        )
        renewed_dep_formula = (
            f'IF(OR({ref("是否需要更新")}<>"是",$CROW<EOMONTH({ref("首次更新月")},{dep_offset})),0,'
            f'IF(MOD(DATEDIF(EOMONTH({ref("首次更新月")},{dep_offset}),$CROW,"m"),{ref("更新后经济耐用月数")})'
            f'<{ref("更新折旧月数")},{ref("更新后每月折旧")},0))'
        )
    else:
        ws.cell(value_rows["折旧开始月"], 2, f'=EOMONTH({ref("投入使用月")},IF({ref("折旧摊销方式")}="折旧",1,0))')
        ws.cell(value_rows["首次更新月"], 2, f'=EOMONTH({ref("投入使用月")},{ref("经济耐用月数")})')
        ws.cell(value_rows["后续更新净支出"], 2, f'=IF({ref("是否需要更新")}="是",{ref("总投资额")}-{ref("残值")},0)')
        renewal_trigger = (
            f'IF({ref("是否需要更新")}<>"是",0,IF($CROW<={ref("投入使用月")},0,'
            f'IF(MOD(DATEDIF({ref("投入使用月")},$CROW,"m"),{ref("经济耐用月数")})=0,1,0)))'
        )
        dep_offset = f'IF({ref("折旧摊销方式")}="折旧",1,0)'
        initial_dep_formula = "0"
        renewed_dep_formula = (
            f'IF($CROW<{ref("折旧开始月")},0,IF({ref("是否需要更新")}="是",'
            f'IF(MOD(DATEDIF({ref("折旧开始月")},$CROW,"m"),{ref("经济耐用月数")})<{ref("折旧月数")},{ref("每月折旧")},0),'
            f'IF(DATEDIF({ref("折旧开始月")},$CROW,"m")<{ref("折旧月数")},{ref("每月折旧")},0)))'
        )

    monthly_start_row = row + 1
    _style_header(ws, monthly_start_row, 1, len(monthly_headers))
    for offset, header in enumerate(monthly_headers, start=1):
        ws.cell(monthly_start_row, offset, header)

    monthly_data_start = monthly_start_row + 1
    for index, month in enumerate(timeline, start=1):
        row_idx = monthly_data_start + index - 1
        ws.cell(row_idx, 1, index)
        ws.cell(row_idx, 2, f"=YEAR(C{row_idx})")
        ws.cell(row_idx, 3, month)
        if asset.asset_kind == "存量":
            trigger = renewal_trigger.replace("ROW", str(row_idx))
            ws.cell(row_idx, 4, "=0")
            ws.cell(row_idx, 6, f'=IF({trigger}=1,{ref("首次更新净支出")},IF({trigger}=2,{ref("后续更新净支出")},0))')
            ws.cell(row_idx, 8, f'=IF({trigger}=1,{ref("原资产残值")},IF({trigger}=2,{ref("更新资产残值")},0))')
            ws.cell(row_idx, 12, f'=IF({trigger}>0,{ref("评估原值")}*{ref("进项税率")},0)')
            ws.cell(row_idx, 14, f'=H{row_idx}*{ref("进项税率")}')
            dep_formula = (
                f'={initial_dep_formula.replace("ROW", str(row_idx))}'
                f'+{renewed_dep_formula.replace("ROW", str(row_idx))}'
            )
            ws.cell(row_idx, 16, dep_formula)
        else:
            trigger = renewal_trigger.replace("ROW", str(row_idx))
            ws.cell(row_idx, 4, f'=IF($C{row_idx}={ref("投入使用月")},{ref("追加资本性支出")},0)')
            ws.cell(row_idx, 6, f'=IF({trigger}=1,{ref("后续更新净支出")},0)')
            ws.cell(row_idx, 8, f'=IF({trigger}=1,{ref("残值")},0)')
            ws.cell(row_idx, 12, f'=IF($C{row_idx}={ref("投入使用月")},D{row_idx}*{ref("进项税率")},0)+IF({trigger}=1,{ref("总投资额")}*{ref("进项税率")},0)')
            ws.cell(row_idx, 14, f'=H{row_idx}*{ref("进项税率")}')
            ws.cell(row_idx, 16, f'={renewed_dep_formula.replace("ROW", str(row_idx))}')
        ws.cell(row_idx, 5, f'=D{row_idx}*R{row_idx}')
        ws.cell(row_idx, 7, f'=F{row_idx}*R{row_idx}')
        ws.cell(row_idx, 9, f'=H{row_idx}*R{row_idx}')
        ws.cell(row_idx, 10, f'=D{row_idx}+F{row_idx}-H{row_idx}')
        ws.cell(row_idx, 11, f'=J{row_idx}*R{row_idx}')
        ws.cell(row_idx, 13, f'=L{row_idx}*R{row_idx}')
        ws.cell(row_idx, 15, f'=N{row_idx}*R{row_idx}')
        ws.cell(
            row_idx,
            17,
            f'=IF(C{row_idx}<=$B$11,0,12*(YEAR(C{row_idx})-YEAR($B$11))+MONTH(C{row_idx})-MONTH($B$11))',
        )
        ws.cell(row_idx, 18, f'=IF(Q{row_idx}=0,1,1/(1+$B$15)^Q{row_idx})')
        ws.cell(row_idx, 19, f'=J{row_idx}*R{row_idx}')
        ws.cell(row_idx, 20, f'=P{row_idx}*R{row_idx}')

    summary_metrics = [
        ("折旧摊销", "T"),
        ("追加资本性支出", "E"),
        ("更新资本性支出", "G"),
        ("残值回收", "I"),
        ("更新金额（净额）", "K"),
        ("进项税", "M"),
        ("销项税", "O"),
        ("资本性支出", None),
    ]
    monthly_end_row = monthly_data_start + len(timeline) - 1
    perpetual_start_row = monthly_data_start
    perpetual_end_row = monthly_end_row
    for metric_row, (metric_name, column_letter) in enumerate(summary_metrics, start=4):
        ws.cell(metric_row, summary_start_col, metric_name)
        for idx, year in enumerate(years, start=1):
            col = get_column_letter(summary_start_col + idx)
            if metric_name == "资本性支出":
                # 资本性支出 = 追加现值 + 更新现值（残值回收是回收，不计入支出）
                formula = (
                    f'=SUMIFS($E${monthly_data_start}:$E${monthly_end_row},'
                    f'$B${monthly_data_start}:$B${monthly_end_row},{col}$3,'
                    f'$C${monthly_data_start}:$C${monthly_end_row},"<="&$B$11)'
                    f'+SUMIFS($G${monthly_data_start}:$G${monthly_end_row},'
                    f'$B${monthly_data_start}:$B${monthly_end_row},{col}$3,'
                    f'$C${monthly_data_start}:$C${monthly_end_row},"<="&$B$11)'
                )
            else:
                formula = (
                    f'=SUMIFS(${column_letter}${monthly_data_start}:${column_letter}${monthly_end_row},'
                    f'$B${monthly_data_start}:$B${monthly_end_row},{col}$3,'
                    f'$C${monthly_data_start}:$C${monthly_end_row},"<="&$B$11)'
                )
            ws.cell(metric_row, summary_start_col + idx, formula)
        if metric_name == "资本性支出":
            formula = (
                f'=(SUMIFS($E${perpetual_start_row}:$E${perpetual_end_row},'
                f'$C${perpetual_start_row}:$C${perpetual_end_row},">"&$B$11,'
                f'$C${perpetual_start_row}:$C${perpetual_end_row},"<="&$B$13)'
                f'+SUMIFS($G${perpetual_start_row}:$G${perpetual_end_row},'
                f'$C${perpetual_start_row}:$C${perpetual_end_row},">"&$B$11,'
                f'$C${perpetual_start_row}:$C${perpetual_end_row},"<="&$B$13))*$B$14'
            )
        else:
            formula = (
                f'=SUMIFS(${column_letter}${perpetual_start_row}:${column_letter}${perpetual_end_row},'
                f'$C${perpetual_start_row}:$C${perpetual_end_row},">"&$B$11,'
                f'$C${perpetual_start_row}:$C${perpetual_end_row},"<="&$B$13)*$B$14'
            )
        ws.cell(metric_row, summary_start_col + len(years) + 1, formula)
        ws.cell(
            metric_row,
            summary_start_col + len(years) + 2,
            "年度合计=∑当年逐月明细；永续期=∑详细预测期后逐月折现现值 × 折现率（年化数）",
        )

    _style_header(ws, 3, summary_start_col, summary_start_col + len(years) + 2)
    _style_header(ws, monthly_start_row, 1, len(monthly_headers))
    ws.freeze_panes = f"A{monthly_data_start}"
    ws.auto_filter.ref = f"A{monthly_start_row}:{get_column_letter(len(monthly_headers))}{monthly_end_row}"
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 14
    for col_letter in ("D", "E", "F", "G", "H", "I", "J", "K", "L", "M", "N", "O", "P", "Q", "R", "S", "T"):
        ws.column_dimensions[col_letter].width = 14
    for col in range(summary_start_col, summary_start_col + len(years) + 3):
        ws.column_dimensions[get_column_letter(col)].width = 14
    for r in range(4, row):
        ws.cell(r, 1).fill = SECTION_FILL
        ws.cell(r, 1).font = Font(bold=True)
        ws.cell(r, 2).fill = INPUT_FILL
        label = ws.cell(r, 1).value
        if label in {"评估基准日", "结束日期", "详细预测期后截止日", "启用时间", "预计投入使用时间", "投入使用月", "预测开始月", "首次更新月", "折旧开始月"}:
            ws.cell(r, 2).number_format = "yyyy-mm-dd"
        elif label in {"折现率", "月折现率", "残值率", "进项税率"}:
            ws.cell(r, 2).number_format = "0.00%"
        elif isinstance(ws.cell(r, 2).value, (int, float, date)):
            if isinstance(ws.cell(r, 2).value, date):
                ws.cell(r, 2).number_format = "yyyy-mm-dd"
            else:
                ws.cell(r, 2).number_format = '#,##0.00;[Red](#,##0.00);-'
    for r in range(monthly_data_start, monthly_end_row + 1):
        for c in range(1, 21):
            ws.cell(r, c).border = TABLE_BORDER
        ws.cell(r, 3).number_format = "yyyy-mm-dd"
        ws.cell(r, 18).number_format = "0.000000"
        for c in range(1, 21):
            if c in (3, 18):
                continue
            ws.cell(r, c).number_format = '#,##0.00;[Red](#,##0.00);-'
    for r in range(4, 4 + len(summary_metrics)):
        for c in range(summary_start_col, summary_start_col + len(years) + 3):
            ws.cell(r, c).border = TABLE_BORDER
        for c in range(summary_start_col + 1, summary_start_col + len(years) + 2):
            ws.cell(r, c).number_format = '#,##0.00;[Red](#,##0.00);-'
        ws.cell(r, summary_start_col + len(years) + 2).alignment = Alignment(wrap_text=True)

    _format_output(ws)


def _write_monthly_detail_sheet(
    ws,
    assets: list[AssetForecast],
    params: ForecastParams,
    *,
    cutoff_end_date: date,
    horizon_end_date: date,
) -> None:
    timeline = _month_timeline(params.valuation_date, horizon_end_date)
    headers = ["类别", "序号", "公司主体", "资产科目", "名称", "费用科目", "情形描述", "指标", "口径"]
    month_start_col = len(headers) + 1
    perpetual_col = month_start_col + len(timeline)
    last_col = perpetual_col
    monthly_rate = params.monthly_discount_rate
    detail_metrics = [
        "折旧摊销",
        "追加资本性支出",
        "更新资本性支出",
        "更新支出进项税",
        "残值回收",
        "残值回收销项税",
    ]

    _clear_sheet(ws)
    _sheet_title(ws, "资产级结果明细（月度）", last_col)
    ws["A2"] = "按月度展开至200年；结束日期前不折现，结束日期后按结束日期逐月折现"
    ws["A2"].font = Font(italic=True, color="808080")
    ws["A2"].alignment = Alignment(wrap_text=True)

    for offset, header in enumerate(headers, start=1):
        ws.cell(3, offset, header)
    for offset, month in enumerate(timeline, start=month_start_col):
        ws.cell(3, offset, month)
    ws.cell(3, perpetual_col, "永续期")
    _style_header(ws, 3, 1, last_col)

    ws.merge_cells(start_row=4, start_column=1, end_row=4, end_column=len(headers))
    ws["A4"] = "折现期（月）"
    ws["A4"].fill = SECTION_FILL
    ws["A4"].font = Font(bold=True)
    ws["A4"].alignment = Alignment(horizontal="center", vertical="center")
    ws["A4"].border = TABLE_BORDER

    ws.merge_cells(start_row=5, start_column=1, end_row=5, end_column=len(headers))
    ws["A5"] = "折现系数"
    ws["A5"].fill = SECTION_FILL
    ws["A5"].font = Font(bold=True)
    ws["A5"].alignment = Alignment(horizontal="center", vertical="center")
    ws["A5"].border = TABLE_BORDER

    for offset, month in enumerate(timeline, start=month_start_col):
        if month <= cutoff_end_date:
            period = None
            factor = None
        else:
            period = _month_index(cutoff_end_date, month)
            factor = 1 / (1 + monthly_rate) ** period if period else 1.0
        period_cell = ws.cell(4, offset, period)
        period_cell.fill = SECTION_FILL
        period_cell.font = Font(bold=True)
        period_cell.alignment = Alignment(horizontal="center", vertical="center")
        if period is not None:
            period_cell.number_format = "0"
        factor_cell = ws.cell(5, offset, factor)
        factor_cell.fill = SECTION_FILL
        factor_cell.font = Font(bold=True)
        factor_cell.alignment = Alignment(horizontal="center", vertical="center")
        if factor is not None:
            factor_cell.number_format = "0.000000"
        ws.cell(3, offset).number_format = "yyyy-mm-dd"

    ws.cell(4, perpetual_col, "")
    ws.cell(5, perpetual_col, "")
    for col in range(1, last_col + 1):
        ws.cell(3, col).border = TABLE_BORDER
    for col in range(month_start_col, last_col + 1):
        ws.cell(4, col).border = TABLE_BORDER
        ws.cell(5, col).border = TABLE_BORDER
    ws.freeze_panes = f"{get_column_letter(month_start_col)}6"
    ws.sheet_view.showGridLines = False

    row = 6
    for asset in assets:
        asset_perpetual = _asset_perpetual_present_values(
            asset,
            cutoff_end_date,
            params.monthly_discount_rate,
            params.discount_rate,
        )
        monthly_values = asset.monthly
        for metric in detail_metrics:
            raw_row = [
                asset.asset_kind,
                asset.asset_id,
                asset.company,
                asset.account,
                asset.name,
                asset.expense_account,
                asset.scenario,
                metric,
                "发生额",
            ]
            pv_row = raw_row[:-1] + ["折现至结束日期现值"]

            for offset, month in enumerate(timeline, start=month_start_col):
                value = monthly_values.get(month)
                amount = _metric_amount(value, metric) if value else 0.0
                raw_row.append(money(amount))
                if month <= cutoff_end_date:
                    pv_row.append(0.0)
                else:
                    period = _month_index(cutoff_end_date, month)
                    factor = 1 / (1 + monthly_rate) ** period if period else 1.0
                    pv_row.append(money(amount * factor))

            raw_row.append("")
            pv_row.append(asset_perpetual[metric])
            for offset, value in enumerate(raw_row, start=1):
                ws.cell(row, offset, value)
            row += 1
            for offset, value in enumerate(pv_row, start=1):
                ws.cell(row, offset, value)
            row += 1

    for r in range(6, row):
        for c in range(1, last_col + 1):
            ws.cell(r, c).border = TABLE_BORDER
        for c in range(month_start_col, last_col + 1):
            ws.cell(r, c).number_format = '#,##0.00;[Red](#,##0.00);-'

    for col in range(1, 10):
        ws.column_dimensions[get_column_letter(col)].width = [10, 12, 14, 14, 16, 14, 18, 18, 14][col - 1]
    for col in range(month_start_col, perpetual_col):
        ws.column_dimensions[get_column_letter(col)].width = 12
    ws.column_dimensions[get_column_letter(perpetual_col)].width = 14


def export_result(
    template_path: str | Path,
    output_path: str | Path,
    result: ForecastResult,
    errors: list[str],
    *,
    include_details: bool = False,
    detail_asset_id: str | None = None,
    detail_asset_kind: str | None = None,
    source_params: ForecastParams | None = None,
    source_stock_assets: list[StockAsset] | None = None,
    source_added_assets: list[AddedAsset] | None = None,
    detail_horizon_years: int = DETAIL_HORIZON_YEARS,
) -> Path:
    template_path = Path(template_path)
    output_path = Path(output_path)
    wb = load_workbook(template_path)
    display_params = result.params
    display_cutoff_end_date = result.params.end_date
    detail_assets = result.assets
    detail_years = _forecast_years(display_params.valuation_date, display_params.end_date)
    if include_details and source_params is not None and source_stock_assets is not None and source_added_assets is not None:
        display_params = ForecastParams(
            valuation_date=source_params.valuation_date,
            end_date=_add_years(source_params.end_date, detail_horizon_years),
            discount_rate=source_params.discount_rate,
            minimum_remaining_years=source_params.minimum_remaining_years,
        )
        detail_assets = [forecast_stock(asset, display_params) for asset in source_stock_assets]
        detail_assets.extend(forecast_added(asset, display_params) for asset in source_added_assets)
        display_cutoff_end_date = source_params.end_date
        detail_years = _forecast_years(display_params.valuation_date, display_params.end_date)
    monthly_assets = detail_assets
    monthly_params = display_params
    monthly_cutoff_end_date = display_cutoff_end_date
    monthly_horizon_end_date = display_params.end_date

    years = _forecast_years(display_params.valuation_date, display_cutoff_end_date)
    annual, perpetual, expense_totals = _output_totals(detail_assets, display_params, display_cutoff_end_date)

    summary = wb[SUMMARY_SHEET]
    _clear_sheet(summary)
    _sheet_title(summary, "总输出表", 2 + len(years))
    summary.append([])
    summary.append(["项目", *years, "永续期"])
    summary_header_row = summary.max_row
    for metric in METRICS:
        summary.append([metric, *(annual[metric, year] for year in years), perpetual[metric]])
    summary.append([])
    summary.append(["费用科目", *years, "永续期"])
    expense_header_row = summary.max_row
    accounts = sorted({asset.expense_account for asset in detail_assets} | set(EXPENSE_ACCOUNTS))
    for account in accounts:
        summary.append([account, *(expense_totals[account, year] for year in years), expense_totals[account, "永续期"]])
    summary.append(["合计", *(annual["折旧摊销", year] for year in years), perpetual["折旧摊销"]])
    expense_total_row = summary.max_row
    summary.append([])
    summary.append(["模型状态", *("" for _ in years), "存在错误" if errors else "OK"])
    status_row = summary.max_row
    _format_output(summary)
    _style_header(summary, summary_header_row, 1, summary.max_column)
    _style_header(summary, expense_header_row, 1, summary.max_column)
    for cell in summary[expense_total_row]:
        cell.border = TOTAL_BORDER
        cell.font = Font(bold=True)
    for cell in summary[status_row]:
        cell.border = TOTAL_BORDER
        if cell.value == "OK":
            cell.fill = OK_FILL
        elif cell.value == "存在错误":
            cell.fill = ERROR_FILL
    if EXPENSE_SHEET in wb.sheetnames:
        del wb[EXPENSE_SHEET]

    if DETAIL_SHEET in wb.sheetnames and LEGACY_DETAIL_SHEET in wb.sheetnames:
        del wb[LEGACY_DETAIL_SHEET]
    elif DETAIL_SHEET not in wb.sheetnames and LEGACY_DETAIL_SHEET in wb.sheetnames:
        wb[LEGACY_DETAIL_SHEET].title = DETAIL_SHEET
    if include_details:
        detail = wb[DETAIL_SHEET] if DETAIL_SHEET in wb.sheetnames else wb.create_sheet(DETAIL_SHEET)
        _clear_sheet(detail)
        detail_metrics = ["折旧摊销", "追加资本性支出", "更新资本性支出", "更新支出进项税", "残值回收", "残值回收销项税"]
        detail_headers = ["类别", "序号", "公司主体", "资产科目", "名称", "费用科目", "情形描述", "指标", "口径", *detail_years, "永续期"]
        _sheet_title(detail, "资产级结果明细", len(detail_headers))
        detail.append([])
        detail.append(detail_headers)
        for asset in detail_assets:
            asset_annual = _asset_yearly_totals(
                asset,
                detail_years,
                display_cutoff_end_date,
                display_params.monthly_discount_rate,
                discounted=False,
            )
            asset_pv_annual = _asset_yearly_totals(
                asset,
                detail_years,
                display_cutoff_end_date,
                display_params.monthly_discount_rate,
                discounted=True,
            )
            asset_perpetual = _asset_perpetual_present_values(
                asset,
                display_cutoff_end_date,
                display_params.monthly_discount_rate,
                display_params.discount_rate,
            )
            for metric in detail_metrics:
                detail.append(
                    [
                        asset.asset_kind,
                        asset.asset_id,
                        asset.company,
                        asset.account,
                        asset.name,
                        asset.expense_account,
                        asset.scenario,
                        metric,
                        "发生额",
                        *(asset_annual[metric, year] for year in detail_years),
                        "",
                    ]
                )
                detail.append(
                    [
                        asset.asset_kind,
                        asset.asset_id,
                        asset.company,
                        asset.account,
                        asset.name,
                        asset.expense_account,
                        asset.scenario,
                        metric,
                        "折现至结束日期现值",
                        *(asset_pv_annual[metric, year] for year in detail_years),
                        asset_perpetual[metric],
                    ]
                )
        _format_output(detail)
        detail.auto_filter.ref = f"A3:{get_column_letter(detail.max_column)}{detail.max_row}"
        detail.column_dimensions["G"].width = 48
        detail.column_dimensions["I"].width = 18
        post_cutoff_start_col = 10 + len(years)
        post_cutoff_end_col = 9 + len(detail_years)
        if post_cutoff_start_col <= post_cutoff_end_col:
            detail.column_dimensions.group(
                get_column_letter(post_cutoff_start_col),
                get_column_letter(post_cutoff_end_col),
                outline_level=1,
                hidden=True,
            )
    else:
        if DETAIL_SHEET in wb.sheetnames:
            del wb[DETAIL_SHEET]
        if LEGACY_DETAIL_SHEET in wb.sheetnames:
            del wb[LEGACY_DETAIL_SHEET]

    checks = wb[CHECKS_SHEET]
    _clear_sheet(checks)
    _sheet_title(checks, "检查结果", 5)
    checks.append([])
    checks.append(["检查项", "实际值", "期望值", "状态", "说明"])
    check_rows = _check_rows(annual, perpetual, years, errors)
    _write_matrix(checks, check_rows)
    _format_output(checks)
    checks.column_dimensions["E"].width = 68
    checks.column_dimensions["A"].width = 30
    for row in range(4, checks.max_row + 1):
        status = checks.cell(row, 4)
        status.fill = OK_FILL if status.value == "OK" else ERROR_FILL

    if include_details:
        if MONTHLY_DETAIL_SHEET in wb.sheetnames:
            del wb[MONTHLY_DETAIL_SHEET]
        monthly = wb.create_sheet(MONTHLY_DETAIL_SHEET)
        _write_monthly_detail_sheet(
            monthly,
            monthly_assets,
            monthly_params,
            cutoff_end_date=monthly_cutoff_end_date,
            horizon_end_date=monthly_horizon_end_date,
        )
    elif MONTHLY_DETAIL_SHEET in wb.sheetnames:
        del wb[MONTHLY_DETAIL_SHEET]

    detail_matches: list[AssetForecast] = []
    if detail_asset_id:
        if detail_asset_id == "random":
            detail_matches = [random.choice(detail_assets)] if detail_assets else []
        else:
            extended_params = None
            if source_params is not None and source_stock_assets is not None and source_added_assets is not None:
                extended_params = ForecastParams(
                    valuation_date=source_params.valuation_date,
                    end_date=_add_years(source_params.end_date, detail_horizon_years),
                    discount_rate=source_params.discount_rate,
                    minimum_remaining_years=source_params.minimum_remaining_years,
                )
                for asset in source_stock_assets:
                    if asset.asset_id == str(detail_asset_id) and (detail_asset_kind is None or detail_asset_kind == "存量"):
                        detail_matches = [forecast_stock(asset, extended_params)]
                        break
                if not detail_matches:
                    for asset in source_added_assets:
                        if asset.asset_id == str(detail_asset_id) and (detail_asset_kind is None or detail_asset_kind == "新增"):
                            detail_matches = [forecast_added(asset, extended_params)]
                            break
            if not detail_matches:
                detail_matches = [
                    asset
                    for asset in detail_assets
                    if asset.asset_id == str(detail_asset_id)
                    and (detail_asset_kind is None or asset.asset_kind == detail_asset_kind)
                ]
    if detail_matches:
        if DETAIL_PROCESS_SHEET in wb.sheetnames:
            detail_ws = wb[DETAIL_PROCESS_SHEET]
            _clear_sheet(detail_ws)
        else:
            detail_ws = wb.create_sheet(DETAIL_PROCESS_SHEET)
        _write_detail_process_sheet(
            detail_ws,
            detail_matches[0],
            display_params,
            cutoff_end_date=display_cutoff_end_date,
            horizon_years=detail_horizon_years,
        )
    elif DETAIL_PROCESS_SHEET in wb.sheetnames:
        del wb[DETAIL_PROCESS_SHEET]

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


def _check_rows(
    annual: dict[tuple[str, int], float],
    perpetual: dict[str, float],
    years: list[int],
    errors: list[str],
) -> list[list[Any]]:
    rows: list[list[Any]] = []
    rows.append(["输入校验", len(errors), 0, "OK" if not errors else "错误", "；".join(errors) if errors else "输入字段符合要求"])
    for year in years:
        expected = annual["追加资本性支出", year] + annual["更新资本性支出", year]
        actual = annual["资本性支出", year]
        rows.append([f"{year} 年资本性支出勾稽", actual, expected, "OK" if abs(actual - expected) < 0.01 else "错误", "资本性支出 = 追加资本性支出 + 更新资本性支出"])
    expected = perpetual["追加资本性支出"] + perpetual["更新资本性支出"]
    rows.append(["永续期资本性支出勾稽", perpetual["资本性支出"], expected, "OK" if abs(perpetual["资本性支出"] - expected) < 0.01 else "错误", "资本性支出 = 追加资本性支出 + 更新资本性支出"])
    return rows


def run_workbook(
    input_path: str | Path,
    output_path: str | Path,
    *,
    include_details: bool = False,
    detail_asset_id: str | None = None,
    detail_asset_kind: str | None = None,
) -> Path:
    params, stock_assets, added_assets, errors = read_input(input_path)
    result = forecast_portfolio(params, stock_assets, added_assets)
    return export_result(
        input_path,
        output_path,
        result,
        errors,
        include_details=include_details,
        detail_asset_id=detail_asset_id,
        detail_asset_kind=detail_asset_kind,
        source_params=params,
        source_stock_assets=stock_assets,
        source_added_assets=added_assets,
    )
