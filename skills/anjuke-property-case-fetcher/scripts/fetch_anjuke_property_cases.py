#!/usr/bin/env python3
"""Capture Anjuke property cases and preserve source evidence."""

from __future__ import annotations

import argparse
import csv
from html import escape
import json
import os
import re
import shutil
import sys
import time
from dataclasses import asdict, dataclass, fields
from pathlib import Path
from typing import Any, Callable, Iterable
from urllib.parse import urljoin

try:
    from playwright.sync_api import TimeoutError as PlaywrightTimeoutError
    from playwright.sync_api import sync_playwright
except Exception:
    PlaywrightTimeoutError = TimeoutError  # type: ignore[assignment]
    sync_playwright = None  # type: ignore[assignment]


def clean(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def nullable(value: Any) -> str | None:
    text = clean(value)
    return text or None


def parse_number(value: Any) -> float | None:
    text = clean(value).replace(",", "")
    match = re.search(r"\d+(?:\.\d+)?", text)
    return float(match.group(0)) if match else None


def parse_yuan(value: Any) -> float | None:
    text = clean(value).replace(",", "")
    number = parse_number(text)
    if number is None:
        return None
    if "亿" in text:
        return number * 100_000_000
    if "万" in text:
        return number * 10_000
    return number


def first_match(text: str, patterns: Iterable[str]) -> str | None:
    for pattern in patterns:
        match = re.search(pattern, text, re.I | re.S)
        if match:
            return nullable(match.group(1))
    return None


def label_value(text: str, labels: Iterable[str]) -> str | None:
    label_pattern = "|".join(re.escape(label) for label in labels)
    return first_match(text, [rf"(?:{label_pattern})\s*[：:]\s*([^\n|；;，,。]{{1,100}})"])


def infer_case_type(url: str, text: str, requested: str) -> str:
    if requested in {"sale", "rent"}:
        return requested
    rent_score = sum(marker in f"{url} {text}" for marker in ("出租", "租金", "元/月", "押一付"))
    sale_score = sum(marker in f"{url} {text}" for marker in ("出售", "总价", "售价", "万元"))
    return "rent" if rent_score > sale_score else "sale"


def verification_required(text: str) -> bool:
    return any(marker in text for marker in ("验证码", "访问过于频繁", "安全验证", "滑块", "人机验证"))


DETAIL_SIGNAL_MARKERS = ("总价", "售价", "参考售价", "报价", "租金", "建筑面积", "房屋单价", "单价", "户型", "楼层", "朝向", "装修", "房源编号", "写字楼", "商铺")
DETAIL_DELAY_SECONDS = 2.0


def is_anjuke_detail_url(value: Any) -> bool:
    from urllib.parse import urlparse

    try:
        parsed = urlparse(str(value or "").strip())
    except ValueError:
        return False
    if parsed.scheme != "https" or not (parsed.hostname or "").lower().endswith("anjuke.com"):
        return False
    path = parsed.path or ""
    if re.search(r"/sp-(?:shou|rent)/", path, re.I) and not re.search(r"/sp-(?:shou|rent)/(?:[^/]+/)*\d+(?:/|$)", path, re.I):
        return False
    return bool(
        re.search(r"/prop/view/[A-Za-z0-9_-]+", path, re.I)
        or re.search(r"/(?:fang5|sale|rent)/[A-Za-z0-9_-]+", path, re.I)
        or re.search(r"/(?:xzl-shou|xzl-zu|sp-shou|sp-zu|sp-rent)/(?:[^/]+/)*\d+(?:/|$)", path, re.I)
        or re.search(r"/\d+(?:/|$)", path)
        or re.search(r"\.html$", path, re.I)
    )


def is_recommendation_url(value: Any) -> bool:
    from urllib.parse import urlparse

    try:
        query = urlparse(str(value or "")).query
    except ValueError:
        return False
    return bool(re.search(r"(?:guessrecommend|recommend|history|similar|related)", query, re.I))


def canonical_detail_url(value: Any) -> str:
    from urllib.parse import urlsplit, urlunsplit

    parsed = urlsplit(str(value or "").strip())
    return urlunsplit((parsed.scheme, parsed.netloc, parsed.path, "", ""))


def is_scoped_listing_url(value: Any) -> bool:
    from urllib.parse import urlparse

    parsed = urlparse(str(value or "").strip())
    if parsed.scheme != "https" or not (parsed.hostname or "").lower().endswith("anjuke.com"):
        return False
    path = (parsed.path or "").rstrip("/")
    return bool(re.fullmatch(r"/(?:xzl-shou|xzl-zu|sp-shou|sp-zu|sp-rent)/[^/]+(?:/[^/]+)*", path, re.I)) and not bool(re.search(r"/\d+$", path))


def is_listing_url(value: Any) -> bool:
    from urllib.parse import urlparse

    parsed = urlparse(str(value or "").strip())
    if parsed.scheme != "https" or not (parsed.hostname or "").lower().endswith("anjuke.com"):
        return False
    path = (parsed.path or "").rstrip("/")
    return bool(re.fullmatch(r"/(?:xzl-shou|xzl-zu|sp-shou|sp-zu|sp-rent)(?:/[^/]+)*", path, re.I)) and not bool(re.search(r"/\d+$", path))


def has_detail_signals(text: str) -> bool:
    normalized = clean(text)
    return sum(marker in normalized for marker in DETAIL_SIGNAL_MARKERS) >= 2


def validate_detail_capture(url: str, text: str) -> None:
    if not is_anjuke_detail_url(url):
        raise ValueError("ANJUKE_DETAIL_ROUTE_INVALID")
    if is_recommendation_url(url):
        raise ValueError("ANJUKE_RECOMMENDATION_PAGE")
    if verification_required(text):
        raise ValueError("ANJUKE_DETAIL_VERIFICATION_REQUIRED")
    if not has_detail_signals(text):
        raise ValueError("ANJUKE_DETAIL_PAGE_NOT_CASE")


def parse_coordinate(value: Any, minimum: float, maximum: float) -> float | None:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if minimum <= number <= maximum else None


def extract_coordinates(value: Any) -> tuple[float | None, float | None]:
    text = str(value or "")
    patterns = [
        (r'''(?:longitude|lng|lon)\s*["':= ]+(-?\d+(?:\.\d+)?)[\s\S]{0,120}?(?:latitude|lat)\s*["':= ]+(-?\d+(?:\.\d+)?)''', "longitude_first"),
        (r'''(?:latitude|lat)\s*["':= ]+(-?\d+(?:\.\d+)?)[\s\S]{0,120}?(?:longitude|lng|lon)\s*["':= ]+(-?\d+(?:\.\d+)?)''', "latitude_first"),
    ]
    for pattern, order in patterns:
        match = re.search(pattern, text, re.I)
        if not match:
            continue
        first, second = match.groups()
        if order == "latitude_first":
            latitude = parse_coordinate(first, -90, 90)
            longitude = parse_coordinate(second, -180, 180)
        else:
            longitude = parse_coordinate(first, -180, 180)
            latitude = parse_coordinate(second, -90, 90)
        if longitude is not None and latitude is not None:
            return longitude, latitude
    return None, None


def ensure_output(path: str | Path) -> Path:
    output = Path(path).expanduser().resolve()
    output.mkdir(parents=True, exist_ok=True)
    (output / "html").mkdir(exist_ok=True)
    (output / "screenshots").mkdir(exist_ok=True)
    return output


def chrome_executables() -> list[str]:
    candidates = [
        "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
        "/Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge",
        os.path.expandvars(r"%PROGRAMFILES%\\Google\\Chrome\\Application\\chrome.exe"),
        os.path.expandvars(r"%LOCALAPPDATA%\\Google\\Chrome\\Application\\chrome.exe"),
        os.path.expandvars(r"%PROGRAMFILES%\\Microsoft\\Edge\\Application\\msedge.exe"),
    ]
    for command in ("google-chrome", "google-chrome-stable", "chromium", "chromium-browser", "msedge"):
        resolved = shutil.which(command)
        if resolved:
            candidates.append(resolved)
    return list(dict.fromkeys(path for path in candidates if path and Path(path).exists()))


def launch_persistent(playwright: Any, user_data_dir: str, headed: bool) -> Any:
    options = {"headless": not headed, "viewport": {"width": 1440, "height": 1000}}
    try:
        return playwright.chromium.launch_persistent_context(user_data_dir, channel="chrome", **options)
    except Exception as channel_error:
        for executable in chrome_executables():
            try:
                return playwright.chromium.launch_persistent_context(user_data_dir, executable_path=executable, **options)
            except Exception:
                continue
        raise RuntimeError("ANJUKE_CHROME_NOT_FOUND: 未找到可用的本机 Chrome，请安装 Chrome 后重试。") from channel_error


def launch_browser(playwright: Any, headed: bool) -> Any:
    options = {"headless": not headed}
    try:
        return playwright.chromium.launch(channel="chrome", **options)
    except Exception as channel_error:
        for executable in chrome_executables():
            try:
                return playwright.chromium.launch(executable_path=executable, **options)
            except Exception:
                continue
        raise RuntimeError("ANJUKE_CHROME_NOT_FOUND: 未找到可用的本机 Chrome，请安装 Chrome 后重试。") from channel_error


@dataclass
class CaseRow:
    input_method: str | None
    case_number: int
    location: str | None
    sale_price_yuan: float | None
    building_area_m2: float | None
    tax_rate: float | None
    floor: str | None
    decoration: str | None
    layout: str | None
    transaction_time: str | None
    orientation: str | None
    building_structure: str | None
    lease_status: str | None
    property_type: str | None
    construction_year: int | None
    source_url: str
    regional_location: str | None
    commercial_activity: str | None
    transport: str | None
    environment: str | None
    bay_depth_ratio: str | None
    case_type: str
    total_price_text: str | None
    sale_unit_price_text: str | None
    rent_total_text: str | None
    rent_unit_price_text: str | None
    rent_pricing_basis: str | None
    payment_terms_text: str | None
    title: str | None
    html_path: str | None
    screenshot_path: str | None
    capture_status: str
    longitude: float | None = None
    latitude: float | None = None


EXPORT_HEADERS = [
    "录入方式", "", "案例序号", "位置", "售价", "建筑面积", "税率", "交易价格(元/㎡)",
    "楼层", "装修", "户型", "交易时间", "朝向", "建筑结构", "租赁情况", "物业类型",
    "建成年份", "案例来源", "区域位置", "商业繁华程度", "交通便捷程度", "环境状况", "开间进深比",
]


def captured_evidence_html(url: str, title: str | None, location: str | None, body: str) -> str:
    return (
        "<!doctype html><meta charset=\"utf-8\"><title>安居客页面快照</title>"
        + f"<h1>{escape(title or '安居客案例')}</h1><p>来源：{escape(url)}</p>"
        + f"<p>位置：{escape(location or '')}</p><pre>{escape(body)}</pre>"
    )


def build_case_row(index: int, url: str, output: Path, requested_type: str, body: str, title: str | None, location: str | None, page_html: str, screenshot_file: Path | None, longitude: float | None = None, latitude: float | None = None) -> CaseRow:
    body = clean(body)
    validate_detail_capture(url, body)
    title = nullable(title)
    location = nullable(location) or title
    if longitude is None or latitude is None:
        longitude, latitude = extract_coordinates(page_html)
    html_file = output / "html" / f"case_{index:03d}.html"
    html_file.write_text(page_html or captured_evidence_html(url, title, location, body), encoding="utf-8")
    case_type = infer_case_type(url, body, requested_type)
    total_price_text = first_match(body, [r"(?:总价|售价)\s*[：:]?\s*([\d,.]+\s*(?:万|万元|亿))", r"([\d,.]+\s*(?:万|万元))"])
    sale_unit_price_text = first_match(body, [r"(?:单价|售价单价)\s*[：:]?\s*([\d,.]+\s*元/(?:㎡|平米|平|m²))", r"([\d,.]+\s*元/(?:㎡|平米|平|m²))"])
    rent_total_text = first_match(body, [r"(?:租金|月租)\s*[：:]?\s*([\d,.]+\s*元/(?:月|天|年))", r"([\d,.]+\s*元/月)"])
    rent_unit_price_text = first_match(body, [r"([\d,.]+\s*元/(?:㎡|平米|平|m²)/(?:天|月))"])
    area_text = label_value(body, ("建筑面积", "面积")) or first_match(body, [r"([\d,.]+\s*(?:㎡|平米|m²|平方米))"])
    floor = label_value(body, ("楼层", "所在楼层")) or first_match(body, [r"((?:低区|中区|高区|\d+层|共\d+层)[^\s，,。]{0,12})"])
    publish = label_value(body, ("发布时间", "更新于", "交易时间")) or first_match(body, [r"(20\d{2}[-年./]\d{1,2}[-月./]\d{1,2})"])
    construction = parse_number(label_value(body, ("建成年份", "竣工年份", "年代")))
    construction_year = int(construction) if construction and construction >= 1800 else None
    property_type = first_match(body, [r"(办公|商铺|住宅|商住楼|写字楼|厂房)"])
    return CaseRow(
        input_method="数据库", case_number=index, location=location,
        sale_price_yuan=parse_yuan(total_price_text) if case_type == "sale" else None,
        building_area_m2=parse_number(area_text), tax_rate=0.09 if case_type == "sale" else None,
        floor=floor, decoration=first_match(body, [r"(精装修|简装|毛坯|豪装|中装)"]),
        layout=label_value(body, ("户型", "格局")), transaction_time=publish,
        orientation=label_value(body, ("朝向",)), building_structure=label_value(body, ("建筑结构",)),
        lease_status=label_value(body, ("租赁情况", "带租约", "出租情况")), property_type=property_type,
        construction_year=construction_year, source_url=url, regional_location=location,
        commercial_activity=None, transport=None, environment=None, bay_depth_ratio=None,
        case_type=case_type, total_price_text=total_price_text, sale_unit_price_text=sale_unit_price_text,
        rent_total_text=rent_total_text, rent_unit_price_text=rent_unit_price_text,
        rent_pricing_basis=first_match(body, [r"(元/(?:㎡|平米|平|m²)/(?:天|月))"]),
        payment_terms_text=first_match(body, [r"((?:押\d+付\d+|面议))"]), title=title,
        html_path=str(html_file), screenshot_path=str(screenshot_file) if screenshot_file else None,
        capture_status="ok",
        longitude=longitude,
        latitude=latitude,
    )


def extract_case(page: Any, index: int, url: str, output: Path, requested_type: str, screenshot: bool, wait_verification: bool, verification_timeout: int) -> CaseRow:
    page.goto(url, wait_until="domcontentloaded", timeout=60_000)
    try:
        page.wait_for_load_state("networkidle", timeout=15_000)
    except PlaywrightTimeoutError:
        pass
    page.wait_for_timeout(1_000)
    body = clean(page.evaluate("() => document.body ? (document.body.innerText || document.body.textContent || '') : ''"))
    blocked = verification_required(body)
    if blocked and wait_verification:
        deadline = time.time() + max(1, verification_timeout)
        while time.time() < deadline and blocked:
            page.wait_for_timeout(3_000)
            body = clean(page.evaluate("() => document.body ? (document.body.innerText || document.body.textContent || '') : ''"))
            blocked = verification_required(body)
    validate_detail_capture(url, body)
    title = nullable(page.evaluate("() => document.querySelector('h1,.title,.house-title,.main-title')?.innerText || document.title || ''"))
    page_html = page.content()
    html_file = output / "html" / f"case_{index:03d}.html"
    html_file.write_text(page_html, encoding="utf-8")
    screenshot_file: Path | None = None
    if screenshot:
        screenshot_file = output / "screenshots" / f"case_{index:03d}_fullpage.png"
        page.screenshot(path=str(screenshot_file), full_page=True)

    case_type = infer_case_type(url, body, requested_type)
    total_price_text = first_match(body, [r"(?:总价|售价)\s*[：:]?\s*([\d,.]+\s*(?:万|万元|亿))", r"([\d,.]+\s*(?:万|万元))"])
    sale_unit_price_text = first_match(body, [r"(?:单价|售价单价)\s*[：:]?\s*([\d,.]+\s*元/(?:㎡|平米|平|m²))", r"([\d,.]+\s*元/(?:㎡|平米|平|m²))"])
    rent_total_text = first_match(body, [r"(?:租金|月租)\s*[：:]?\s*([\d,.]+\s*元/(?:月|天|年))", r"([\d,.]+\s*元/月)"])
    rent_unit_price_text = first_match(body, [r"([\d,.]+\s*元/(?:㎡|平米|平|m²)/(?:天|月))"])
    area_text = label_value(body, ("建筑面积", "面积")) or first_match(body, [r"([\d,.]+\s*(?:㎡|平米|m²|平方米))"])
    floor = label_value(body, ("楼层", "所在楼层")) or first_match(body, [r"((?:低区|中区|高区|\d+层|共\d+层)[^\s，,。]{0,12})"])
    publish = label_value(body, ("发布时间", "更新于", "交易时间")) or first_match(body, [r"(20\d{2}[-年./]\d{1,2}[-月./]\d{1,2})"])
    construction = parse_number(label_value(body, ("建成年份", "竣工年份", "年代")))
    construction_year = int(construction) if construction and construction >= 1800 else None
    location = nullable(page.evaluate("() => document.querySelector('.address,.addr,[class*=address],[class*=addr]')?.innerText || ''")) or title
    property_type = first_match(body, [r"(办公|商铺|住宅|商住楼|写字楼|厂房)"])
    transaction_time = publish
    return CaseRow(
        input_method="数据库",
        case_number=index,
        location=location,
        sale_price_yuan=parse_yuan(total_price_text) if case_type == "sale" else None,
        building_area_m2=parse_number(area_text),
        tax_rate=0.09 if case_type == "sale" else None,
        floor=floor,
        decoration=first_match(body, [r"(精装修|简装|毛坯|豪装|中装)"]),
        layout=label_value(body, ("户型", "格局")),
        transaction_time=transaction_time,
        orientation=label_value(body, ("朝向",)),
        building_structure=label_value(body, ("建筑结构",)),
        lease_status=label_value(body, ("租赁情况", "带租约", "出租情况")),
        property_type=property_type,
        construction_year=construction_year,
        source_url=url,
        regional_location=location,
        commercial_activity=None,
        transport=None,
        environment=None,
        bay_depth_ratio=None,
        case_type=case_type,
        total_price_text=total_price_text,
        sale_unit_price_text=sale_unit_price_text,
        rent_total_text=rent_total_text,
        rent_unit_price_text=rent_unit_price_text,
        rent_pricing_basis=first_match(body, [r"(元/(?:㎡|平米|平|m²)/(?:天|月))"]),
        payment_terms_text=first_match(body, [r"((?:押\d+付\d+|面议))"]),
        title=title,
        html_path=str(html_file),
        screenshot_path=str(screenshot_file) if screenshot_file else None,
        capture_status="ok",
    )


def extract_captured_case(capture: dict[str, Any], index: int, output: Path, requested_type: str) -> CaseRow:
    url = str(capture.get("url") or "").strip()
    title = nullable(capture.get("title"))
    location = nullable(capture.get("location"))
    body = clean(capture.get("text"))
    page_html = str(capture.get("html") or "")
    longitude = parse_coordinate(capture.get("longitude"), -180, 180)
    latitude = parse_coordinate(capture.get("latitude"), -90, 90)
    return build_case_row(index, url, output, requested_type, body, title, location, page_html, None, longitude, latitude)


def collect_detail_urls(page: Any, list_url: str, keyword: str, max_cases: int, url_pattern: str) -> list[str]:
    if not is_listing_url(list_url):
        raise ValueError("ANJUKE_LIST_PAGE_REQUIRED")
    page.goto(list_url, wait_until="domcontentloaded", timeout=60_000)
    try:
        page.wait_for_load_state("networkidle", timeout=15_000)
    except PlaywrightTimeoutError:
        pass
    for _ in range(12):
        page.evaluate("window.scrollTo(0, document.body.scrollHeight)")
        page.wait_for_timeout(650)
    page.evaluate("window.scrollTo(0, 0)")
    page.wait_for_timeout(500)
    anchors = page.locator("a[href]")
    links = []
    for index in range(anchors.count()):
        anchor = anchors.nth(index)
        href = anchor.get_attribute("href") or ""
        text = clean(anchor.inner_text())
        card_text = anchor.evaluate("""a => { let node=a; let value=''; for (let i=0; node && i<6; i++, node=node.parentElement) { const text=(node.innerText||node.textContent||'').trim(); if (text.length>value.length && text.length<=2400) value=text; } return value; }""")
        links.append({"href": href, "text": text, "cardText": card_text})
    detail_re = re.compile(url_pattern)
    result: list[str] = []
    seen: set[str] = set()
    for item in links:
        href = str(item.get("href") or "").split("#", 1)[0]
        text = clean(item.get("text"))
        if keyword and keyword not in f"{text} {href}":
            # Listing cards often keep the property name in a parent node.
            text = clean(item.get("cardText"))
        if not href or href in seen or is_recommendation_url(href) or not is_anjuke_detail_url(href) or not detail_re.search(href):
            continue
        if keyword and keyword not in f"{text} {href}":
            continue
        href = canonical_detail_url(href)
        seen.add(href)
        result.append(href)
        if len(result) >= max_cases:
            break
    return result


def excel_values(row: CaseRow) -> list[Any]:
    return [
        row.input_method, None, row.case_number, row.location, row.sale_price_yuan, row.building_area_m2,
        row.tax_rate, None, row.floor, row.decoration, row.layout, row.transaction_time, row.orientation,
        row.building_structure, row.lease_status, row.property_type, row.construction_year, row.source_url,
        row.regional_location, row.commercial_activity, row.transport, row.environment, row.bay_depth_ratio,
    ]


def write_csv(rows: list[CaseRow], output: Path) -> Path:
    target = output / "cases.csv"
    keys = [field.name for field in fields(CaseRow)]
    with target.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=keys)
        writer.writeheader()
        writer.writerows(asdict(row) for row in rows)
    return target


def write_excel(rows: list[CaseRow], output: Path) -> Path:
    try:
        from openpyxl import Workbook
    except Exception as exc:
        raise RuntimeError("ANJUKE_OPENPYXL_REQUIRED") from exc
    target = output / "cases.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "安居客案例"
    sheet.append(EXPORT_HEADERS)
    for cell in sheet[1]:
        cell.font = cell.font.copy(bold=True)
    for row_index, row in enumerate(rows, 2):
        values = excel_values(row)
        for column, value in enumerate(values, 1):
            sheet.cell(row_index, column, value)
        sheet.cell(row_index, 8).value = f"=ROUND(E{row_index}/F{row_index}/(1+G{row_index}),0)" if row.sale_price_yuan and row.building_area_m2 else None
        sheet.cell(row_index, 12).number_format = "@"
        if row.transaction_time:
            sheet.cell(row_index, 12).value = str(row.transaction_time)
        label = clean(row.title) or clean(row.location) or "安居客案例"
        safe_url = row.source_url.replace('"', '""')
        safe_label = label.replace('"', '""')
        sheet.cell(row_index, 18).value = f'=HYPERLINK("{safe_url}","安居客-{safe_label}")'
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:W{max(1, len(rows) + 1)}"
    for column in (4, 18, 19, 20, 21, 22, 23):
        sheet.column_dimensions[chr(64 + column)].width = 18
    workbook.save(target)
    return target


def write_json(rows: list[CaseRow], output: Path) -> Path:
    target = output / "cases.json"
    target.write_text(json.dumps([asdict(row) for row in rows], ensure_ascii=False, indent=2), encoding="utf-8")
    return target


def write_result_pages(rows: list[CaseRow], output: Path) -> tuple[Path, Path]:
    result_path = output / "result.html"
    map_path = output / "map.html"
    table_rows = []
    for index, row in enumerate(rows, 1):
        coordinate = "已定位" if row.longitude is not None and row.latitude is not None else "无坐标"
        table_rows.append(
            "<tr>"
            f"<td>{index}</td><td>{escape(row.title or '')}</td><td>{escape(row.location or '')}</td>"
            f"<td>{escape(row.total_price_text or row.rent_total_text or '')}</td>"
            f"<td>{escape(str(row.building_area_m2 or ''))}</td><td>{escape(coordinate)}</td>"
            f"<td><a href=\"{escape(row.source_url, quote=True)}\" target=\"_blank\" rel=\"noreferrer\">来源</a>"
            f" &middot; <a href=\"map.html#case-{index}\">地图</a></td></tr>"
        )
    result_path.write_text(
        "<!doctype html><meta charset=\"utf-8\"><title>安居客案例结果</title>"
        "<style>body{font:14px system-ui;margin:24px;color:#1f2937}h1{margin-bottom:6px}"
        "table{border-collapse:collapse;width:100%;margin-top:18px}th,td{border:1px solid #d1d5db;padding:8px;text-align:left}"
        "th{background:#f3f4f6}a{color:#2563eb}</style>"
        "<h1>安居客案例结果</h1>"
        f"<p>共 {len(rows)} 条。数据来源为逐个打开的安居客详情页，验证页和通用页面未纳入结果。</p>"
        "<table><thead><tr><th>序号</th><th>标题</th><th>位置</th><th>价格</th><th>面积</th><th>坐标</th><th>链接</th></tr></thead>"
        f"<tbody>{''.join(table_rows)}</tbody></table>\n",
        encoding="utf-8",
    )
    map_data = json.dumps(
        [
            {"id": index, "title": row.title or "安居客案例", "location": row.location or "", "url": row.source_url,
             "longitude": row.longitude, "latitude": row.latitude}
            for index, row in enumerate(rows, 1)
        ], ensure_ascii=False,
    ).replace("<", "\\u003c")
    map_script_prefix = "<script>const cases=" + map_data + ";const htmlEscape=value=>String(value??'').replace(/[&<>\"']/g,character=>({'&':'&amp;','<':'&lt;','>':'&gt;','\"':'&quot;',\"'\":'&#39;'}[character]));const list=document.getElementById('cases');"
    map_path.write_text(
        "<!doctype html><meta charset=\"utf-8\"><title>安居客案例地图</title>"
        "<link rel=\"stylesheet\" href=\"https://unpkg.com/leaflet@1.9.4/dist/leaflet.css\">"
        "<style>body{margin:0;font:14px system-ui;color:#1f2937;display:flex;height:100vh}"
        "aside{width:320px;overflow:auto;padding:16px;background:#f8fafc}#map{flex:1}"
        ".case{display:block;width:100%;text-align:left;border:1px solid #d1d5db;background:white;padding:10px;margin:0 0 8px;cursor:pointer}"
        ".case:hover{background:#eff6ff}.muted{color:#6b7280}</style>"
        "<aside><h2>安居客案例</h2><p class=\"muted\">点击清单定位标记；没有坐标的案例仍保留在结果表格。</p><div id=\"cases\"></div></aside><div id=\"map\"></div>"
        + map_script_prefix
        + "for(const item of cases){const button=document.createElement('button');button.className='case';button.textContent=item.id+' · '+item.title+' '+item.location;button.dataset.id=item.id;list.appendChild(button);}"
        + "const located=cases.filter(item=>Number.isFinite(item.latitude)&&Number.isFinite(item.longitude));"
        + "function noMap(){document.getElementById('map').innerHTML='<div style=\"padding:24px\">当前结果没有可用坐标，仍可在左侧查看案例清单。</div>'; }"
        + "if(!located.length){noMap();}else{const script=document.createElement('script');script.src='https://unpkg.com/leaflet@1.9.4/dist/leaflet.js';script.onload=()=>{const map=L.map('map');L.tileLayer('https://{s}.tile.openstreetmap.org/{z}/{x}/{y}.png',{attribution:'© OpenStreetMap'}).addTo(map);const markers=new Map();for(const item of located){const marker=L.marker([item.latitude,item.longitude]).addTo(map).bindPopup('<b>'+htmlEscape(item.id+' · '+item.title)+'</b><br>'+htmlEscape(item.location));markers.set(String(item.id),marker);}map.fitBounds(L.featureGroup([...markers.values()]).getBounds().pad(0.2));for(const button of document.querySelectorAll('.case'))button.onclick=()=>{const marker=markers.get(button.dataset.id);if(marker){map.setView(marker.getLatLng(),15);marker.openPopup();}};};script.onerror=noMap;document.head.appendChild(script);}"
        + "</script>\n",
        encoding="utf-8",
    )
    return result_path, map_path


def run_captured_request(request: dict[str, Any], captured_pages: list[dict[str, Any]], progress: Callable[[dict[str, Any]], None] | None = None) -> dict[str, Any]:
    def emit(payload: dict[str, Any]) -> None:
        if progress:
            progress(payload)

    output = ensure_output(request.get("outputDirectory") or request.get("out") or "")
    case_type = clean(request.get("caseType", request.get("case_type", "auto"))) or "auto"
    max_cases = max(1, min(100, int(request.get("maxCases", request.get("max_cases", 10)) or 10)))
    if bool(request.get("screenshot", False)):
        return {"ok": False, "errorCode": "ANJUKE_CURRENT_TAB_SCREENSHOT_UNSUPPORTED", "reason": "当前浏览器标签页模式暂不支持详情页全页截图，请取消勾选“保存全页截图”后重试。", "security": {"credentialsReturned": False}}
    pages: list[dict[str, Any]] = []
    seen: set[str] = set()
    for page in captured_pages:
        url = str(page.get("url") or "").strip()
        if not url or url in seen or not page.get("text"):
            continue
        seen.add(url)
        pages.append(page)
        if len(pages) >= max_cases:
            break
    if not pages:
        return {"ok": False, "errorCode": "ANJUKE_NO_DETAIL_URLS", "reason": "当前安居客页面没有找到可读取的详情案例；请确认范围已加载完成。", "security": {"credentialsReturned": False}}
    emit({"phase": "capturing", "percent": 10, "message": "正在读取当前浏览器标签页中的安居客详情…", "fetched": len(pages), "written": 0})
    rows: list[CaseRow] = []
    skipped_invalid = max(0, int(request.get("skippedInvalidCount", request.get("skipped_invalid_count", 0)) or 0))
    try:
        for index, page in enumerate(pages, 1):
            try:
                row = extract_captured_case(page, index, output, case_type)
            except (ValueError, TypeError) as exc:
                if str(exc) in {"ANJUKE_DETAIL_ROUTE_INVALID", "ANJUKE_RECOMMENDATION_PAGE", "ANJUKE_DETAIL_VERIFICATION_REQUIRED", "ANJUKE_DETAIL_PAGE_NOT_CASE"}:
                    skipped_invalid += 1
                    continue
                raise
            rows.append(row)
            emit({"phase": "capturing", "percent": 10 + round(index / len(pages) * 80), "message": f"已完成 {index}/{len(pages)} 条详情页证据保存。", "fetched": len(pages), "written": len(rows), "current": row.title or row.source_url})
    except Exception as exc:
        return {"ok": False, "errorCode": "ANJUKE_CAPTURED_PAGE_PARSE_FAILED", "reason": str(exc)[:300], "results": [asdict(row) for row in rows], "security": {"credentialsReturned": False}}
    if not rows:
        return {"ok": False, "errorCode": "ANJUKE_NO_VALID_DETAIL_CASES", "reason": "读取到的页面均不是有效安居客详情案例，未生成案例文件。", "skippedInvalidCount": skipped_invalid, "security": {"credentialsReturned": False}}
    emit({"phase": "writing", "percent": 94, "message": "正在生成 CSV、JSON、结果表格和地图…", "fetched": len(pages), "written": len(rows)})
    csv_path = write_csv(rows, output)
    json_path = write_json(rows, output)
    result_html_path, map_path = write_result_pages(rows, output)
    try:
        excel_path = write_excel(rows, output)
    except Exception as exc:
        return {"ok": False, "errorCode": str(exc), "reason": "Excel 生成失败，已保留 CSV/JSON 和当前页证据。", "csvPath": str(csv_path), "jsonPath": str(json_path), "results": [asdict(row) for row in rows], "security": {"credentialsReturned": False}}
    return {
        "ok": True, "errorCode": "", "reason": "安居客详情案例抓取、证据归档、结果表格和地图输出完成。",
        "caseCount": len(rows), "blockedVerificationCount": 0, "skippedInvalidCount": skipped_invalid,
        "outputDirectory": str(output), "csvPath": str(csv_path), "jsonPath": str(json_path),
        "excelPath": str(excel_path), "htmlDirectory": str(output / "html"), "resultHtmlPath": str(result_html_path),
        "mapPath": str(map_path), "screenshotDirectory": "",
        "results": [asdict(row) for row in rows], "security": {"credentialsReturned": False},
    }


def run_request(request: dict[str, Any], progress: Callable[[dict[str, Any]], None] | None = None) -> dict[str, Any]:
    def emit(payload: dict[str, Any]) -> None:
        if progress:
            progress(payload)

    output = ensure_output(request.get("outputDirectory") or request.get("out") or "")
    captured_pages = request.get("capturedPages", request.get("captured_pages"))
    if isinstance(captured_pages, list):
        return run_captured_request(request, captured_pages, progress)
    list_urls = [str(value).strip() for value in request.get("listUrls", request.get("list_urls", [])) if str(value).strip()]
    detail_urls = [str(value).strip() for value in request.get("detailUrls", request.get("detail_urls", [])) if str(value).strip()]
    keyword = clean(request.get("keyword"))
    case_type = clean(request.get("caseType", request.get("case_type", "auto"))) or "auto"
    max_cases = max(1, min(100, int(request.get("maxCases", request.get("max_cases", 10)) or 10)))
    url_pattern = str(request.get("urlPattern", request.get("url_pattern", r"anjuke\.com/.*/\d+/?")))
    screenshot = bool(request.get("screenshot", False))
    headed = bool(request.get("headed", False))
    wait_verification = bool(request.get("waitVerification", request.get("wait_verification", False)))
    verification_timeout = max(1, min(900, int(request.get("verificationTimeout", 300) or 300)))
    if not list_urls and not detail_urls:
        return {"ok": False, "errorCode": "ANJUKE_SOURCE_URL_REQUIRED", "reason": "请填写列表页 URL 或至少一个详情页 URL。", "security": {"credentialsReturned": False}}
    if sync_playwright is None:
        return {"ok": False, "errorCode": "ANJUKE_PLAYWRIGHT_REQUIRED", "reason": "本机缺少 Playwright，请安装后重试。", "security": {"credentialsReturned": False}}
    emit({"phase": "opening", "percent": 5, "message": "正在启动安居客受控浏览器…", "fetched": 0, "written": 0})
    rows: list[CaseRow] = []
    skipped_invalid = 0
    try:
        with sync_playwright() as playwright:
            browser = None
            user_data_dir = clean(request.get("userDataDir", request.get("user_data_dir", "")))
            if user_data_dir:
                context = launch_persistent(playwright, str(Path(user_data_dir).expanduser().resolve()), headed)
            else:
                browser = launch_browser(playwright, headed)
                context = browser.new_context(viewport={"width": 1440, "height": 1000})
            page = context.new_page()
            urls = detail_urls[:]
            try:
                for list_url in list_urls:
                    urls.extend(collect_detail_urls(page, list_url, keyword, max_cases, url_pattern))
            except ValueError as exc:
                context.close()
                if browser:
                    browser.close()
                code = str(exc) or "ANJUKE_LIST_CAPTURE_FAILED"
                return {"ok": False, "errorCode": code, "reason": "来源网址不是带明确区域条件的列表页，已阻止抓取。", "security": {"credentialsReturned": False}}
            deduped: list[str] = []
            seen: set[str] = set()
            for url in urls:
                canonical = urljoin("https://hz.sydc.anjuke.com", url)
                if canonical not in seen:
                    seen.add(canonical)
                    deduped.append(canonical)
                if len(deduped) >= max_cases:
                    break
            if not deduped:
                context.close()
                if browser:
                    browser.close()
                return {"ok": False, "errorCode": "ANJUKE_NO_DETAIL_URLS", "reason": "没有找到符合条件的安居客详情页。", "security": {"credentialsReturned": False}}
            for index, url in enumerate(deduped, 1):
                if index > 1:
                    time.sleep(DETAIL_DELAY_SECONDS)
                try:
                    row = extract_case(page, index, url, output, case_type, screenshot, wait_verification, verification_timeout)
                except (ValueError, TypeError) as exc:
                    if str(exc) in {"ANJUKE_DETAIL_ROUTE_INVALID", "ANJUKE_RECOMMENDATION_PAGE", "ANJUKE_DETAIL_VERIFICATION_REQUIRED", "ANJUKE_DETAIL_PAGE_NOT_CASE"}:
                        skipped_invalid += 1
                        continue
                    raise
                rows.append(row)
                emit({"phase": "capturing", "percent": 10 + round(index / len(deduped) * 80), "message": f"已完成 {index}/{len(deduped)} 条详情页证据保存。", "fetched": len(deduped), "written": len(rows), "current": row.title or url})
            context.close()
            if browser:
                browser.close()
    except Exception as exc:
        return {"ok": False, "errorCode": "ANJUKE_CAPTURE_FAILED", "reason": str(exc)[:300], "results": [asdict(row) for row in rows], "security": {"credentialsReturned": False}}
    if not rows:
        return {"ok": False, "errorCode": "ANJUKE_NO_VALID_DETAIL_CASES", "reason": "读取到的页面均不是有效安居客详情案例，未生成案例文件。", "skippedInvalidCount": skipped_invalid, "security": {"credentialsReturned": False}}
    emit({"phase": "writing", "percent": 94, "message": "正在生成 CSV、JSON、结果表格和地图…", "fetched": len(deduped), "written": len(rows)})
    csv_path = write_csv(rows, output)
    json_path = write_json(rows, output)
    result_html_path, map_path = write_result_pages(rows, output)
    try:
        excel_path = write_excel(rows, output)
    except Exception as exc:
        return {"ok": False, "errorCode": str(exc), "reason": "Excel 生成失败，已保留 CSV/JSON 和原始证据。", "csvPath": str(csv_path), "jsonPath": str(json_path), "results": [asdict(row) for row in rows], "security": {"credentialsReturned": False}}
    return {
        "ok": True,
        "errorCode": "",
        "reason": "安居客详情案例抓取、证据归档、结果表格和地图输出完成。",
        "caseCount": len(rows), "skippedInvalidCount": skipped_invalid,
        "blockedVerificationCount": 0,
        "outputDirectory": str(output),
        "csvPath": str(csv_path),
        "jsonPath": str(json_path),
        "excelPath": str(excel_path),
        "htmlDirectory": str(output / "html"), "resultHtmlPath": str(result_html_path), "mapPath": str(map_path),
        "screenshotDirectory": str(output / "screenshots") if screenshot else "",
        "results": [asdict(row) for row in rows],
        "security": {"credentialsReturned": False},
    }


def open_only(url: str, user_data_dir: str) -> int:
    if sync_playwright is None:
        print("Playwright unavailable", file=sys.stderr)
        return 1
    with sync_playwright() as playwright:
        context = launch_persistent(playwright, str(Path(user_data_dir).expanduser().resolve()), True)
        page = context.new_page()
        page.goto(url, wait_until="domcontentloaded", timeout=60_000)
        while not page.is_closed():
            page.wait_for_timeout(1_000)
        context.close()
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description="抓取安居客物业出售/租赁案例并输出证据和 A-W Excel")
    parser.add_argument("--request-json", default="")
    parser.add_argument("--list-url", action="append", default=[])
    parser.add_argument("--detail-url", action="append", default=[])
    parser.add_argument("--keyword", default="")
    parser.add_argument("--case-type", choices=["auto", "sale", "rent"], default="auto")
    parser.add_argument("--out", default="")
    parser.add_argument("--max-cases", type=int, default=10)
    parser.add_argument("--url-pattern", default=r"anjuke\.com/.*/\d+/?")
    parser.add_argument("--headed", action="store_true")
    parser.add_argument("--screenshot", action="store_true")
    parser.add_argument("--wait-verification", action="store_true")
    parser.add_argument("--verification-timeout", type=int, default=300)
    parser.add_argument("--user-data-dir", default="")
    parser.add_argument("--open-only", action="store_true")
    parser.add_argument("--url", default="")
    args = parser.parse_args()
    if args.open_only:
        if not args.url or not args.user_data_dir:
            return 1
        return open_only(args.url, args.user_data_dir)
    if args.request_json:
        request = json.loads(args.request_json)
    else:
        request = {
            "listUrls": args.list_url, "detailUrls": args.detail_url, "keyword": args.keyword,
            "caseType": args.case_type, "outputDirectory": args.out, "maxCases": args.max_cases,
            "urlPattern": args.url_pattern, "headed": args.headed, "screenshot": args.screenshot,
            "waitVerification": args.wait_verification, "verificationTimeout": args.verification_timeout,
            "userDataDir": args.user_data_dir,
        }
    result = run_request(request, lambda payload: print(f"TY_ANJUKE_PROGRESS:{json.dumps(payload, ensure_ascii=False)}", flush=True))
    print(f"TY_ANJUKE_RESULT:{json.dumps(result, ensure_ascii=False)}", flush=True)
    return 0 if result.get("ok") else 1


if __name__ == "__main__":
    raise SystemExit(main())
